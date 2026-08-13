from __future__ import annotations

import csv
from hashlib import sha256
import os
from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtCore import QItemSelectionModel, QPoint, QTimer, Qt
from PySide6.QtTest import QSignalSpy, QTest
from PySide6.QtWidgets import (
    QApplication,
    QFileDialog,
    QLabel,
    QLineEdit,
    QMessageBox,
    QSplitter,
    QTableWidget,
)
from openpyxl import load_workbook
import pytest

from test_spd_decap_distribution import (
    _direct_scenario,
    _scenario_with_routing_asset,
    _shared_chain_scenario,
    _with_initial_rails,
)
from spd_decap_pi.distribution import (
    DISTRIBUTION_VIA_PROJECTION_POLICY,
    DistributionDiagnostic,
    DistributionDistanceMode,
    DistributionPlanStatus,
    apply_distribution_plan,
    compute_distribution_plan,
)
from spd_decap_pi.distribution_workbook import (
    DISTRIBUTION_METADATA_TITLE,
    DISTRIBUTION_TOLERANCE_SEMANTICS,
    load_distribution_targets,
)
from spd_decap_pi.evaluation import preflight_evaluation_connectivity
from spd_decap_pi.gui.main_window import MainWindow, _job_compute_distribution
from spd_decap_pi.routing_obstacles import SignalTraceAvoidancePolicy
from spd_decap_pi.scenario import (
    DecapConnectionKind,
    DecapPadState,
    ScenarioDecapConnection,
    ScenarioSpec,
    SharedPadClusterState,
)
from spd_decap_pi.scenario_io import load_scenario
from spd_decap_pi.spreadsheet_export import write_distribution_workbook


def _application() -> QApplication:
    return QApplication.instance() or QApplication([])


def _window_with_scenario(scenario: ScenarioSpec) -> MainWindow:
    window = MainWindow()
    window._scenario = scenario
    window._attachments = {}
    window._dirty = False
    window._invalidate_evaluation("Distribution GUI fixture")
    window._refresh_all()
    return window


def _rail_row(window: MainWindow, rail_id: str) -> int:
    for row in range(window.distribution_table.rowCount()):
        item = window.distribution_table.item(row, 0)
        if str(item.data(Qt.ItemDataRole.UserRole)).casefold() == rail_id.casefold():
            return row
    raise AssertionError(f"rail {rail_id!r} is absent from Distribution table")


def _set_target(
    window: MainWindow,
    rail_id: str,
    value: int | str,
    *,
    model_id: str = "M1",
) -> None:
    row = _rail_row(window, rail_id)
    column = next(
        index
        for index in range(window.distribution_table.columnCount())
        if window.distribution_table.horizontalHeaderItem(index).text()
        == f"{model_id}\nTarget"
    )
    window.distribution_table.item(row, column).setText(str(value))


def _targets(window: MainWindow) -> dict[tuple[str, str], int]:
    return dict(window._distribution_targets)


def _set_tolerance(
    window: MainWindow,
    rail_id: str,
    value: float | str,
    *,
    model_id: str = "M1",
) -> None:
    row = _rail_row(window, rail_id)
    column = next(
        index
        for index in range(window.distribution_table.columnCount())
        if window.distribution_table.horizontalHeaderItem(index).text()
        == f"{model_id}\nTolerance (%)"
    )
    window.distribution_table.item(row, column).setText(str(value))


def test_distribution_tab_matches_the_target_matrix_and_resizable_sections() -> None:
    application = _application()
    scenario = _direct_scenario(
        (
            ("C1", 0.0, ("R1", "R2")),
            ("C2", 10.0, ("R1", "R2")),
            ("C3", 20.0, ("R1", "R2")),
        )
    )
    window = _window_with_scenario(scenario)
    try:
        window.resize(1200, 700)
        window.side_tabs.setCurrentIndex(3)
        window.show()
        application.processEvents()
        assert [
            window.side_tabs.tabText(index)
            for index in range(window.side_tabs.count())
        ] == ["Selection", "Evaluation", "AI Assist", "De-cap Distribution"]
        splitter = window.findChild(QSplitter, "distributionSectionSplitter")
        assert splitter is not None
        assert splitter.count() == 2
        assert not splitter.childrenCollapsible()
        assert tuple(
            splitter.widget(index).minimumHeight() for index in range(2)
        ) == (280, 180)
        targets_size, results_size = splitter.sizes()
        assert targets_size >= 280
        assert results_size >= 180
        for button in (
            window.import_distribution_targets_button,
            window.calculate_distribution_button,
            window.apply_distribution_button,
            window.export_distribution_csv_button,
            window.save_distribution_button,
        ):
            assert button.isVisible()
            assert button.width() >= button.sizeHint().width()

        assert window.distribution_table.rowCount() == 3
        plane_note = window.findChild(QLabel, "alternatePwrPlaneRoutingNote")
        assert plane_note is not None
        assert "Vertical VIA projection" in plane_note.text()
        assert "MLO transition/short-span evidence does not block" in plane_note.text()
        assert "exact target-plane copper" in plane_note.text()
        assert "does not certify" in plane_note.toolTip()
        assert window.distribution_table.selectionMode() == (
            window.distribution_table.SelectionMode.ExtendedSelection
        )
        assert window.distribution_table.columnCount() == 5
        assert not window.source_board_checkbox.isEnabled()
        assert window.board_assignment_view_label.text() == (
            "Board: Current (matches source)"
        )
        assert window.distribution_table.horizontalHeaderItem(0).text() == "PWR NET"
        assert window.distribution_table.horizontalHeaderItem(1).text() == (
            "M1\nPresent"
        )
        assert window.distribution_table.horizontalHeaderItem(2).text() == (
            "M1\nTarget"
        )
        assert window.distribution_table.horizontalHeaderItem(3).text() == (
            "M1\nTolerance (%)"
        )
        assert window.distribution_table.horizontalHeaderItem(4).text() == (
            "M1\nAssignment Failed"
        )
        row = _rail_row(window, "R1")
        present = window.distribution_table.item(row, 1)
        target = window.distribution_table.item(row, 2)
        tolerance = window.distribution_table.item(row, 3)
        assignment_failed = window.distribution_table.item(row, 4)
        assert present.text() == target.text() == "3"
        assert tolerance.text() == "0"
        assert not present.flags() & Qt.ItemFlag.ItemIsEditable
        assert target.flags() & Qt.ItemFlag.ItemIsEditable
        assert tolerance.flags() & Qt.ItemFlag.ItemIsEditable
        assert not assignment_failed.flags() & Qt.ItemFlag.ItemIsEditable
        assert "Unfulfilled receiver demand" in assignment_failed.toolTip()
        assert window.distribution_distance_combo.itemData(0) == "NEAREST"
        assert window.distribution_distance_combo.itemData(1) == "FARTHEST"
        assert not window.distribution_protect_signal_routing_checkbox.isChecked()
        assert not window.distribution_trace_clearance_edit.isEnabled()
        scope_note = window.findChild(QLabel, "distributionSignalRoutingScopeNote")
        assert scope_note is not None
        assert "SIGNAL Trace" in scope_note.text()
        assert not window.calculate_distribution_button.isEnabled()
        assert window.distribution_validation_label.text() == (
            "M1: Donor 0 | Receiver 0 | Balance +0"
        )
        assert "every Target equals Present" in window.distribution_summary.toPlainText()
    finally:
        window._dirty = False
        window.close()
        application.processEvents()


def test_detached_distribution_editor_routes_target_tolerance_and_invalid_state(
) -> None:
    application = _application()
    window = _window_with_scenario(
        _direct_scenario((("C1", 0.0, ("R1", "R2")),), rail_ids=("R1", "R2"))
    )
    try:
        window._show_distribution_window()
        dialog = window._distribution_window
        assert dialog is not None
        assert dialog.table.editTriggers() & QTableWidget.EditTrigger.DoubleClicked
        assert dialog.table.editTriggers() & QTableWidget.EditTrigger.EditKeyPressed
        assert (
            dialog.table.selectionMode()
            == QTableWidget.SelectionMode.ExtendedSelection
        )
        target = dialog.table.item(0, 2)
        key = tuple(target.data(Qt.ItemDataRole.UserRole))
        main_row = _rail_row(window, key[0])
        for column in (0, 1, 4):
            assert not (
                dialog.table.item(0, column).flags()
                & Qt.ItemFlag.ItemIsEditable
            )
        for column in (2, 3):
            detached_item = dialog.table.item(0, column)
            main_item = window.distribution_table.item(main_row, column)
            assert detached_item.flags() == main_item.flags()
            assert detached_item.toolTip() == main_item.toolTip()
            assert detached_item.textAlignment() == main_item.textAlignment()
            assert detached_item.flags() & Qt.ItemFlag.ItemIsEditable
        window._distribution_plan = object()
        window._distribution_preview_scenario = object()
        calls: list[int] = []
        original = window._distribution_targets_edited

        def record_edit() -> None:
            calls.append(1)
            original()

        window._distribution_targets_edited = record_edit
        dialog.table.editItem(target)
        editor = dialog.table.findChild(QLineEdit)
        assert editor is not None
        editor.selectAll()
        editor.setText("7")
        QTest.keyClick(editor, Qt.Key.Key_Return)
        application.processEvents()
        assert calls == [1]
        assert window._distribution_targets[key] == 7
        assert window._distribution_plan is None
        assert window._distribution_preview_scenario is None
        assert (
            window.distribution_table.item(_rail_row(window, key[0]), 2).text()
            == "7"
        )
        target = dialog.table.item(0, 2)
        assert target.data(Qt.ItemDataRole.UserRole) == key
        tolerance = dialog.table.item(0, 3)
        tolerance_key = tuple(tolerance.data(Qt.ItemDataRole.UserRole))
        assert tolerance.data(Qt.ItemDataRole.UserRole + 1) == "tolerance"
        assert tolerance.flags() & Qt.ItemFlag.ItemIsEditable
        dialog.table.editItem(tolerance)
        application.processEvents()
        editor = QApplication.focusWidget()
        assert editor is not None
        assert isinstance(editor, QLineEdit)
        assert editor.isVisible()
        editor.selectAll()
        editor.setText("1.5")
        QTest.keyClick(editor, Qt.Key.Key_Return)
        application.processEvents()
        tolerance = dialog.table.item(0, 3)
        assert tolerance.text() == "1.5"
        assert window._distribution_tolerances[tolerance_key] == 1.5
        assert calls == [1, 1]
        assert (
            window.distribution_table.item(
                _rail_row(window, tolerance_key[0]), 3
            ).text()
            == "1.5"
        )
        target = dialog.table.item(0, 2)
        target.setText("-1")
        application.processEvents()
        target = dialog.table.item(0, 2)
        main_target = window.distribution_table.item(_rail_row(window, key[0]), 2)
        assert target.text() == main_target.text() == "-1"
        assert target.background() == main_target.background()
        assert key in window._distribution_invalid_cells
        target.setText("7")
        application.processEvents()
        tolerance = dialog.table.item(0, 3)
        tolerance.setText("nan")
        application.processEvents()
        tolerance = dialog.table.item(0, 3)
        main_tolerance = window.distribution_table.item(
            _rail_row(window, tolerance_key[0]), 3
        )
        assert tolerance.text() == main_tolerance.text() == "nan"
        assert tolerance.background() == main_tolerance.background()
        assert tolerance_key in window._distribution_invalid_tolerance_cells
        assert not window.calculate_distribution_button.isEnabled()
    finally:
        window._dirty = False
        window.close()
        application.processEvents()


def test_detached_distribution_same_field_bulk_edit_commits_once(
    monkeypatch,
) -> None:
    application = _application()
    window = _window_with_scenario(
        _direct_scenario(
            (
                ("C1", 0.0, ("R1", "R2")),
                ("C2", 10.0, ("R1", "R2")),
            ),
            rail_ids=("R1", "R2"),
        )
    )
    try:
        window._show_distribution_window()
        dialog = window._distribution_window
        assert dialog is not None
        first = dialog.table.item(0, 2)
        second = dialog.table.item(1, 2)
        mixed_tolerance = dialog.table.item(1, 3)
        first_key = tuple(first.data(Qt.ItemDataRole.UserRole))
        second_key = tuple(second.data(Qt.ItemDataRole.UserRole))
        for item in (first, second, mixed_tolerance):
            item.setSelected(True)
        dialog.table.setCurrentItem(first, QItemSelectionModel.SelectionFlag.NoUpdate)
        window._distribution_plan = object()
        window._distribution_preview_scenario = object()
        calls = 0
        original = window._distribution_targets_edited

        def record_edit() -> None:
            nonlocal calls
            calls += 1
            original()

        monkeypatch.setattr(window, "_distribution_targets_edited", record_edit)
        dialog.table.editItem(first)
        editor = dialog.table.findChild(QLineEdit)
        assert editor is not None
        editor.selectAll()
        editor.setText("4")
        QTest.keyClick(editor, Qt.Key.Key_Return)
        application.processEvents()

        assert window._distribution_targets[first_key] == 4
        assert window._distribution_targets[second_key] == 4
        assert window._distribution_tolerances[second_key] == 0.0
        assert calls == 1
        assert window._distribution_plan is None
        assert window._distribution_preview_scenario is None
        assert dialog.table.item(0, 2).text() == "4"
        assert dialog.table.item(1, 2).text() == "4"
        assert dialog.table.item(1, 3).text() == "0"
    finally:
        window._dirty = False
        window.close()
        application.processEvents()


def test_detached_unchanged_source_text_fills_selected_tolerance_once(
    monkeypatch,
) -> None:
    application = _application()
    window = _window_with_scenario(
        _direct_scenario(
            (
                ("C1", 0.0, ("R1", "R2")),
                ("C2", 10.0, ("R1", "R2")),
            ),
            rail_ids=("R1", "R2"),
        )
    )
    try:
        _set_tolerance(window, "R1", 1.5)
        window._show_distribution_window()
        dialog = window._distribution_window
        assert dialog is not None
        first = dialog.table.item(_rail_row(window, "R1"), 3)
        second = dialog.table.item(_rail_row(window, "R2"), 3)
        first.setSelected(True)
        second.setSelected(True)
        dialog.table.setCurrentItem(first, QItemSelectionModel.SelectionFlag.NoUpdate)
        calls = 0
        original = window._distribution_targets_edited

        def record_edit() -> None:
            nonlocal calls
            calls += 1
            original()

        monkeypatch.setattr(window, "_distribution_targets_edited", record_edit)
        dialog.table.editItem(first)
        editor = dialog.table.findChild(QLineEdit)
        assert editor is not None
        assert editor.text() == "1.5"
        QTest.keyClick(editor, Qt.Key.Key_Return)
        application.processEvents()

        assert window._distribution_tolerances[("R1", "M1")] == 1.5
        assert window._distribution_tolerances[("R2", "M1")] == 1.5
        assert calls == 1
        assert dialog.table.item(_rail_row(window, "R2"), 3).text() == "1.5"
    finally:
        window._dirty = False
        window.close()
        application.processEvents()


def test_detached_invalid_same_field_bulk_edit_marks_every_cell_once(
    monkeypatch,
) -> None:
    application = _application()
    window = _window_with_scenario(
        _direct_scenario(
            (
                ("C1", 0.0, ("R1", "R2")),
                ("C2", 10.0, ("R1", "R2")),
            ),
            rail_ids=("R1", "R2"),
        )
    )
    try:
        window._show_distribution_window()
        dialog = window._distribution_window
        assert dialog is not None
        first = dialog.table.item(0, 2)
        second = dialog.table.item(1, 2)
        keys = (
            tuple(first.data(Qt.ItemDataRole.UserRole)),
            tuple(second.data(Qt.ItemDataRole.UserRole)),
        )
        first.setSelected(True)
        second.setSelected(True)
        dialog.table.setCurrentItem(first, QItemSelectionModel.SelectionFlag.NoUpdate)
        calls = 0
        original = window._distribution_targets_edited

        def record_edit() -> None:
            nonlocal calls
            calls += 1
            original()

        monkeypatch.setattr(window, "_distribution_targets_edited", record_edit)
        dialog.table.editItem(first)
        editor = dialog.table.findChild(QLineEdit)
        assert editor is not None
        editor.selectAll()
        editor.setText("-1")
        QTest.keyClick(editor, Qt.Key.Key_Return)
        application.processEvents()

        assert calls == 1
        assert set(keys).issubset(window._distribution_invalid_cells)
        for row, key in enumerate(keys):
            detached = dialog.table.item(row, 2)
            main = window.distribution_table.item(_rail_row(window, key[0]), 2)
            assert detached.text() == main.text() == "-1"
            assert detached.background() == main.background()
        assert not window.calculate_distribution_button.isEnabled()
    finally:
        window._dirty = False
        window.close()
        application.processEvents()


def test_detached_distribution_busy_edit_is_disabled() -> None:
    application = _application()
    window = _window_with_scenario(
        _direct_scenario((("C1", 0.0, ("R1", "R2")),), rail_ids=("R1", "R2"))
    )
    try:
        window._show_distribution_window()
        dialog = window._distribution_window
        assert dialog is not None
        window._worker = object()
        window._sync_distribution_window()
        item = dialog.table.item(0, 2)
        key = tuple(item.data(Qt.ItemDataRole.UserRole))
        before = item.text()
        value = window._distribution_targets[key]
        assert not dialog.table.isEnabled()
        dialog.table.setCurrentItem(item)
        dialog.table.setFocus()
        QTest.keyClick(dialog.table, Qt.Key.Key_9)
        application.processEvents()
        assert item.text() == before
        assert window._distribution_targets[key] == value
    finally:
        window._worker = None
        window._set_busy(False)
        window._dirty = False
        window.close()
        application.processEvents()


def test_queued_detached_commit_is_discarded_after_busy_transition() -> None:
    application = _application()
    window = _window_with_scenario(
        _direct_scenario(
            (("C1", 0.0, ("R1", "R2")),),
            rail_ids=("R1", "R2"),
        )
    )
    try:
        window._show_distribution_window()
        dialog = window._distribution_window
        assert dialog is not None
        item = dialog.table.item(0, 2)
        key = tuple(item.data(Qt.ItemDataRole.UserRole))
        before = window._distribution_targets[key]
        payload = ("target", "9", (key,), dialog.matrix_revision)
        QTimer.singleShot(0, lambda: dialog.detachedBatchCommitted.emit(payload))

        window._worker = object()
        window._sync_distribution_window()
        application.processEvents()

        assert not dialog.table.isEnabled()
        assert window._distribution_targets[key] == before
    finally:
        window._worker = None
        window._set_busy(False)
        window._dirty = False
        window.close()
        application.processEvents()


def test_distribution_table_double_click_opens_detached_window_and_exports_template(
    tmp_path: Path, monkeypatch
) -> None:
    application = _application()
    scenario = _direct_scenario(
        (("C1", 0.0, ("R1", "R2")),),
        rail_ids=("R1", "R2"),
        bump_x={"R2": 0.0},
    )
    window = _window_with_scenario(scenario)
    try:
        window.side_tabs.setCurrentIndex(3)
        window.show()
        application.processEvents()
        target = window.distribution_table.item(_rail_row(window, "R1"), 2)
        window.distribution_table.itemDoubleClicked.emit(target)
        application.processEvents()

        dialog = window._distribution_window
        assert dialog is not None
        assert dialog.isVisible()
        assert dialog.table.rowCount() == window.distribution_table.rowCount()
        assert dialog.export_template_button.isEnabled()

        path = tmp_path / "targets-template.xlsx"
        monkeypatch.setattr(
            QFileDialog,
            "getSaveFileName",
            lambda *_args, **_kwargs: (str(path), "Excel workbook (*.xlsx)"),
        )
        dialog.export_template_button.click()
        workbook = load_workbook(path, data_only=False)
        try:
            assert workbook.sheetnames == [
                "Decap Changes",
                "PWR NET Distribution Targets",
            ]
            assert workbook["PWR NET Distribution Targets"]["A1"].value == "PWR NET"
        finally:
            workbook.close()
        imported = load_distribution_targets(
            path,
            rail_ids=("R1", "R2"),
            model_ids=("M1",),
            current_present=dict(window._distribution_present_counts),
            current_source_sha256=scenario.source.sha256,
            current_design_fingerprint=scenario.design_fingerprint,
        )
        assert imported.targets == window._distribution_targets
        assert imported.via_projection_policy == DISTRIBUTION_VIA_PROJECTION_POLICY

        # Switching away from a custom policy clears stale input and the
        # template must omit the penalty so it round-trips as MIN_GAPS.
        custom_index = window.distribution_optimization_combo.findData(
            "BALANCED_CUSTOM"
        )
        min_gaps_index = window.distribution_optimization_combo.findData("MIN_GAPS")
        window.distribution_optimization_combo.setCurrentIndex(custom_index)
        window.distribution_gap_penalty_edit.setText("1234.5")
        window.distribution_optimization_combo.setCurrentIndex(min_gaps_index)
        assert window.distribution_gap_penalty_edit.text() == ""
        dialog.export_template_button.click()
        imported_min_gaps = load_distribution_targets(
            path,
            rail_ids=("R1", "R2"),
            model_ids=("M1",),
            current_present=dict(window._distribution_present_counts),
            current_source_sha256=scenario.source.sha256,
            current_design_fingerprint=scenario.design_fingerprint,
        )
        assert imported_min_gaps.optimization_policy == "MIN_GAPS"
        assert imported_min_gaps.effective_gap_penalty_um is None
        assert imported_min_gaps.via_projection_policy == DISTRIBUTION_VIA_PROJECTION_POLICY
    finally:
        window._dirty = False
        window.close()
        application.processEvents()


def test_distribution_balance_strip_stays_compact_and_logs_shortage_narrative() -> None:
    application = _application()
    window = _window_with_scenario(
        _direct_scenario(
            (
                ("C1", 0.0, ("R1", "R2")),
                ("C2", 10.0, ("R1", "R2")),
                ("C3", 20.0, ("R1", "R2")),
            ),
            rail_ids=("R1", "R2"),
        )
    )
    try:
        _set_target(window, "R1", 1)
        _set_target(window, "R2", 4)

        strip = window.distribution_validation_label
        assert not strip.wordWrap()
        assert strip.text() == "M1: Donor 2 | Receiver 4 | Balance -2"
        assert "Invalid" not in strip.text()
        assert "short" not in strip.text().casefold()
        assert "M1 is short by 2" in window.distribution_summary.toPlainText()
    finally:
        window._dirty = False
        window.close()
        application.processEvents()


def test_physical_landing_retarget_has_no_obsolete_column_warning() -> None:
    """The compact balance strip remains free of superseded column wording."""

    application = _application()
    scenario = _direct_scenario(
        (("C1", 5.0, ("R1", "R2")),), rail_ids=("R1", "R2")
    )
    window = _window_with_scenario(scenario)
    try:
        _set_target(window, "R1", 0)
        _set_target(window, "R2", 1)
        plan = compute_distribution_plan(
            scenario,
            _targets(window),
            DistributionDistanceMode.NEAREST,
        )
        assert not any(
            item.code == "DISTRIBUTION_VIA_COLUMN_EVIDENCE_MISSING"
            for item in plan.diagnostics
        )

        window._accept_distribution_plan(plan)

        compact = window.distribution_validation_label.text()
        assert compact == "M1: Donor 1 | Receiver 1 | Balance +0"
        assert "Via-column" not in compact
        assert "reopen" not in compact.casefold()
        assert "Via-column" not in window.distribution_summary.toPlainText()
        assert (
            f"VIA projection: {DISTRIBUTION_VIA_PROJECTION_POLICY}"
            in window.distribution_summary.toPlainText()
        )
    finally:
        window._dirty = False
        window.close()
        application.processEvents()


def test_detached_distribution_controls_follow_main_worker_busy_state() -> None:
    application = _application()
    source = _direct_scenario(
        (("C1", 0.0, ("R1", "R2")),),
        rail_ids=("R1", "R2"),
    )
    distributed = source.model_copy(
        update={
            "decaps": (
                source.decaps[0].model_copy(
                    update={"current_net": "V2", "current_rail_id": "R2"}
                ),
            )
        }
    )
    window = _window_with_scenario(distributed)
    try:
        window._show_distribution_window()
        dialog = window._distribution_window
        assert dialog is not None

        window._worker = object()  # type: ignore[assignment]
        window._set_busy(True)
        assert not dialog.import_targets_button.isEnabled()
        assert not dialog.export_template_button.isEnabled()
        assert not dialog.original_board_checkbox.isEnabled()
        assert not window.source_board_checkbox.isEnabled()

        window._worker = None
        window._set_busy(False)
        assert dialog.import_targets_button.isEnabled()
        assert dialog.export_template_button.isEnabled()
        assert dialog.original_board_checkbox.isEnabled()
        assert window.source_board_checkbox.isEnabled()
    finally:
        window._worker = None
        window._dirty = False
        window.close()
        application.processEvents()


def test_detached_distribution_window_closes_with_main_window() -> None:
    application = _application()
    window = _window_with_scenario(
        _direct_scenario(
            (("C1", 0.0, ("R1", "R2")),),
            rail_ids=("R1", "R2"),
        )
    )
    try:
        window.show()
        window._show_distribution_window()
        dialog = window._distribution_window
        assert dialog is not None
        application.processEvents()
        assert dialog.isVisible()

        window.close()
        application.processEvents()
        assert not dialog.isVisible()
    finally:
        window._dirty = False
        dialog = window._distribution_window
        if dialog is not None:
            dialog.close()
        window.close()
        application.processEvents()


def test_distribution_worker_result_is_discarded_when_targets_change(
    monkeypatch,
) -> None:
    application = _application()
    window = _window_with_scenario(
        _direct_scenario(
            (("C1", 0.0, ("R1", "R2")), ("C2", 10.0, ("R1", "R2"))),
            rail_ids=("R1", "R2"),
        )
    )
    accepted: list[object] = []
    try:
        expected = window._distribution_request_fingerprint()
        _set_target(window, "R1", 1)
        monkeypatch.setattr(
            window, "_accept_distribution_plan", lambda result: accepted.append(result)
        )

        window._accept_distribution_plan_if_current(object(), expected)

        assert accepted == []
        assert "Discarded stale De-cap Distribution preview" == window.status_text.text()
    finally:
        window._dirty = False
        window.close()
        application.processEvents()


def test_routing_option_and_clearance_are_request_inputs_and_fail_closed_without_asset(
    monkeypatch,
) -> None:
    application = _application()
    window = _window_with_scenario(
        _direct_scenario(
            (("C1", 0.0, ("R1", "R2")),), rail_ids=("R1", "R2")
        )
    )
    try:
        baseline = window._distribution_request_fingerprint()
        legacy_payload = repr(
            (
                window._scenario.design_fingerprint,
                window._scenario.revision,
                tuple(
                    sorted(
                        (
                            str(rail_id).casefold(),
                            str(model_id).casefold(),
                            int(value),
                        )
                        for (rail_id, model_id), value in (
                            window._distribution_targets.items()
                        )
                    )
                ),
                tuple(
                    sorted(
                        (
                            str(rail_id).casefold(),
                            str(model_id).casefold(),
                            float(value),
                        )
                        for (rail_id, model_id), value in (
                            window._distribution_tolerances.items()
                        )
                    )
                    ),
                    window.distribution_distance_combo.currentData(),
                    window.distribution_optimization_combo.currentData(),
                    window.distribution_gap_penalty_edit.text().strip(),
                )
        ).encode("utf-8")
        assert baseline == sha256(legacy_payload).hexdigest()
        assert window._distribution_workbook_contract() == (
            5,
            {
                "Signal Routing Protection": "OFF",
                "Tolerance Semantics": DISTRIBUTION_TOLERANCE_SEMANTICS,
            },
        )
        window.distribution_protect_signal_routing_checkbox.setChecked(True)
        enabled_without_clearance = window._distribution_request_fingerprint()
        assert enabled_without_clearance != baseline
        assert window.distribution_trace_clearance_edit.isEnabled()
        assert "Enter Trace-to-via clearance" in window.distribution_summary.toPlainText()

        window.distribution_trace_clearance_edit.setText("12.5")
        with_clearance = window._distribution_request_fingerprint()
        assert with_clearance != enabled_without_clearance
        original_policy = window._distribution_routing_policy
        monkeypatch.setattr(
            window,
            "_distribution_routing_policy",
            lambda: SignalTraceAvoidancePolicy(
                enabled=True,
                clearance_um=12.5,
                policy_version="SIGNAL_NET_ONLY_V_NEXT",
            ),
        )
        assert window._distribution_request_fingerprint() != with_clearance
        monkeypatch.setattr(window, "_distribution_routing_policy", original_policy)
        assert not window.calculate_distribution_button.isEnabled()
        assert "no immutable signal-routing asset" in (
            window.distribution_summary.toPlainText()
        )

        window.distribution_protect_signal_routing_checkbox.setChecked(False)
        assert window._distribution_routing_policy().enabled is False
        assert not window.distribution_trace_clearance_edit.isEnabled()
    finally:
        window._dirty = False
        window.close()
        application.processEvents()


def test_routing_asset_identity_affects_requests_only_when_protection_is_on() -> None:
    application = _application()
    scenario, _attachments, _plane = _scenario_with_routing_asset(trace_y_um=0.0)
    reference = scenario.routing_obstacle_asset
    assert reference is not None
    drifted = scenario.model_copy(
        update={
            "routing_obstacle_asset": reference.model_copy(
                update={"content_sha256": "0" * 64}
            )
        }
    )
    assert drifted.design_fingerprint == scenario.design_fingerprint
    window = _window_with_scenario(scenario)
    try:
        off_original = window._distribution_request_fingerprint()
        window._scenario = drifted
        off_drifted = window._distribution_request_fingerprint()
        assert off_drifted == off_original

        window._scenario = scenario
        window.distribution_protect_signal_routing_checkbox.setChecked(True)
        window.distribution_trace_clearance_edit.setText("0")
        on_original = window._distribution_request_fingerprint()
        window._scenario = drifted
        on_drifted = window._distribution_request_fingerprint()
        assert on_drifted != on_original
    finally:
        window._dirty = False
        window.close()
        application.processEvents()


def test_distribution_worker_result_is_discarded_when_tolerance_or_order_changes(
    monkeypatch,
) -> None:
    application = _application()
    window = _window_with_scenario(
        _direct_scenario(
            (("C1", 0.0, ("R1", "R2")), ("C2", 10.0, ("R1", "R2"))),
            rail_ids=("R1", "R2"),
        )
    )
    accepted: list[object] = []
    try:
        monkeypatch.setattr(
            window, "_accept_distribution_plan", lambda result: accepted.append(result)
        )
        expected_tolerance = window._distribution_request_fingerprint()
        _set_tolerance(window, "R1", 25)
        window._accept_distribution_plan_if_current(object(), expected_tolerance)

        expected_order = window._distribution_request_fingerprint()
        window.distribution_distance_combo.setCurrentIndex(1)
        window._accept_distribution_plan_if_current(object(), expected_order)

        assert accepted == []
        assert "Discarded stale De-cap Distribution preview" == window.status_text.text()
    finally:
        window._dirty = False
        window.close()
        application.processEvents()


def test_detached_distribution_import_applies_targets_without_mutating_scenario(
    tmp_path: Path, monkeypatch
) -> None:
    application = _application()
    scenario = _direct_scenario(
        (("C1", 0.0, ("R1", "R2")), ("C2", 10.0, ("R1", "R2"))),
        rail_ids=("R1", "R2"),
        bump_x={"R2": 0.0},
    )
    path = tmp_path / "edited-targets.xlsx"
    write_distribution_workbook(
        path,
        (),
        ("PWR NET", "M1\nPresent", "M1\nTarget", "M1\nTolerance (%)"),
        (("V1 (R1)", 2, 1, 0), ("V2 (R2)", 0, 1, 0)),
        metadata={
            "Format Version": 2,
            "Source SPD SHA-256": scenario.source.sha256,
            "Distance Mode": "NEAREST",
        },
    )
    window = _window_with_scenario(scenario)
    try:
        window._distribution_plan = object()
        window._distribution_preview_scenario = scenario
        original_fingerprint = scenario.design_fingerprint
        window._show_distribution_window()
        dialog = window._distribution_window
        assert dialog is not None
        monkeypatch.setattr(
            QFileDialog,
            "getOpenFileName",
            lambda *_args, **_kwargs: (str(path), "Excel workbook (*.xlsx)"),
        )
        dialog.import_targets_button.click()
        application.processEvents()

        assert window._distribution_targets[("R1", "M1")] == 1
        assert window._distribution_targets[("R2", "M1")] == 1
        assert window._distribution_plan is None
        assert window._distribution_preview_scenario is None
        assert window.scenario is not None
        assert window.scenario.design_fingerprint == original_fingerprint
        assert not window._dirty
    finally:
        window._dirty = False
        window.close()
        application.processEvents()


def test_board_distribution_toggle_is_display_only() -> None:
    application = _application()
    scenario = _direct_scenario(
        (("C1", 0.0, ("R1", "R2")), ("C2", 10.0, ("R1", "R2"))),
        rail_ids=("R1", "R2"),
        bump_x={"R2": 0.0},
    )
    distributed = ScenarioSpec.model_validate(
        {
            **scenario.model_dump(mode="python"),
            "decaps": [
                scenario.decaps[0].model_copy(
                    update={"current_net": "V2", "current_rail_id": "R2"}
                ),
                scenario.decaps[1],
            ],
            "revision": scenario.revision + 1,
        }
    )
    window = _window_with_scenario(distributed)
    try:
        fingerprint = distributed.design_fingerprint
        revision = distributed.revision
        window._show_distribution_window()
        dialog = window._distribution_window
        assert dialog is not None
        window.board.set_selected_refdes(("C1",))
        window.board._view_box.setRange(
            xRange=(-25.0, 25.0), yRange=(-10.0, 10.0), padding=0.0
        )
        application.processEvents()
        view_before = window.board._view_box.viewRange()
        assert window.source_board_checkbox.isEnabled()
        assert window.board_assignment_view_label.text() == (
            "Board: Current / distributed"
        )
        window.source_board_checkbox.click()
        view_after = window.board._view_box.viewRange()
        assert all(
            abs(before - after) < 1.0e-9
            for before_axis, after_axis in zip(view_before, view_after, strict=True)
            for before, after in zip(before_axis, after_axis, strict=True)
        )
        assert dialog.original_board_checkbox.isChecked()
        assert window.board._records[0].current_net == distributed.decaps[0].source_net
        assert window.board.selected_refdes == ("C1",)
        assert window.board_assignment_view_label.text() == (
            "Board: Source SPD (read-only)"
        )
        assert window.selection_table.item(0, 2).text() == "V1"
        assert window.scenario is not None
        assert window.scenario.design_fingerprint == fingerprint
        assert window.scenario.revision == revision
        assert not window._dirty

        window.search_mode.setCurrentText("PWR NET")
        window.search_edit.setText("V1")
        window.apply_search()
        assert window.board.selected_refdes == ("C1", "C2")
        window._show_decap_context_menu(("C1",), QPoint())
        assert "read-only" in window.status_text.text()

        dialog.original_board_checkbox.click()
        assert not window.source_board_checkbox.isChecked()
        assert window.board._records[0].current_net == "V2"
        window.search_edit.setText("V1")
        window.apply_search()
        assert window.board.selected_refdes == ("C2",)

        cached_source_records = window._source_board_records
        edited_current = distributed.model_copy(
            update={
                "decaps": (
                    distributed.decaps[0],
                    distributed.decaps[1].model_copy(update={"enabled": False}),
                ),
                "revision": distributed.revision + 1,
            }
        )
        window._scenario = edited_current
        window._refresh_all()
        assert window._source_board_records is cached_source_records
        window.source_board_checkbox.click()
        assert all(record.current_net == "V1" for record in window.board._records)
        assert all(record.enabled for record in window.board._records)
    finally:
        window._dirty = False
        window.close()
        application.processEvents()


def test_board_original_toggle_restores_shared_chain_isolation_gap_display() -> None:
    application = _application()
    source = _shared_chain_scenario()
    distributed = apply_distribution_plan(
        source,
        compute_distribution_plan(source, {("R1", "M1"): 1, ("R2", "M1"): 2}),
    )
    window = _window_with_scenario(distributed)
    try:
        current_gap = next(item for item in distributed.decaps if item.refdes == "D1")
        assert not current_gap.enabled
        assert current_gap.pad_state == DecapPadState.ISOLATION_GAP

        window.board.set_selected_refdes(("D1",))
        window._show_original_distribution_board(True)
        original_gap = next(
            item for item in window.board._records if item.refdes == "D1"
        )
        assert original_gap.enabled
        assert original_gap.current_net == current_gap.source_net
        assert window.board._companion_keys == set()
        assert "read-only" in window.selection_summary.text().casefold()
        assert window.selection_table.item(0, 6).text() == "Read-only source view"

        window._show_original_distribution_board(False)
        restored_current = next(
            item for item in window.board._records if item.refdes == "D1"
        )
        assert not restored_current.enabled
    finally:
        window._dirty = False
        window.close()
        application.processEvents()


def test_detached_distribution_window_clears_stale_document_and_exposes_routing_scope() -> None:
    application = _application()
    source = _direct_scenario(
        (("C1", 0.0, ("R1", "R2")),),
        rail_ids=("R1", "R2"),
    )
    distributed = source.model_copy(
        update={
            "decaps": (
                source.decaps[0].model_copy(
                    update={"current_net": "V2", "current_rail_id": "R2"}
                ),
            )
        }
    )
    window = _window_with_scenario(distributed)
    try:
        window._show_distribution_window()
        dialog = window._distribution_window
        assert dialog is not None
        assert "VIA STACK CHANGE REQUIRED" in dialog.alternate_plane_note.text()
        assert "plane artwork remains unchanged" in dialog.alternate_plane_note.text()
        assert "does not prove" in dialog.alternate_plane_note.toolTip()
        assert dialog.table.accessibleName() == "Distribution target matrix"
        assert dialog.import_targets_button.accessibleName() == "Import distribution XLSX"
        assert dialog.export_template_button.accessibleName() == "Export distribution XLSX template"
        assert dialog.original_board_checkbox.accessibleName() == (
            "Show source SPD board assignments"
        )
        window.source_board_checkbox.click()
        assert window.source_board_checkbox.isChecked()

        window._reset_document_view_state()
        assert not dialog.isVisible()
        assert dialog.table.rowCount() == 0
        assert not dialog.import_targets_button.isEnabled()
        assert not dialog.export_template_button.isEnabled()
        assert not dialog.original_board_checkbox.isChecked()
        assert not window.source_board_checkbox.isChecked()
        assert not window.source_board_checkbox.isEnabled()
        assert window.board_assignment_view_label.text() == "Board: No document"
    finally:
        window._dirty = False
        window.close()
        application.processEvents()


def test_legacy_target_import_refreshes_present_requires_distance_and_invalidates_preview(
    tmp_path: Path,
    monkeypatch,
) -> None:
    application = _application()
    scenario = _direct_scenario(
        (
            ("C1", 0.0, ("R1", "R2")),
            ("C2", 10.0, ("R1", "R2")),
            ("C3", 20.0, ("R1", "R2")),
        ),
        rail_ids=("R1", "R2"),
        bump_x={"R2": 0.0},
    )
    path = tmp_path / "legacy-targets.xlsx"
    write_distribution_workbook(
        path,
        (),
        (
            "PWR NET",
            "M1\nPresent",
            "M1\nTarget",
            "M1\nTolerance (%)",
            "M1\nActual Delta",
        ),
        (
            ("V1 (R1)", 1, 1, 0, 0),
            ("V2 (R2)", 0, 2, 0, 2),
        ),
    )
    window = _window_with_scenario(scenario)
    try:
        window._distribution_plan = object()
        window._distribution_preview_scenario = scenario
        window._distribution_power_projection = object()
        window._distribution_export_rows = (object(),)
        monkeypatch.setattr(
            QFileDialog,
            "getOpenFileName",
            lambda *_args, **_kwargs: (str(path), "Excel workbook (*.xlsx)"),
        )

        window.import_distribution_targets_button.click()

        assert window._distribution_plan is None
        assert window._distribution_preview_scenario is None
        assert window._distribution_power_projection is None
        assert window._distribution_export_rows == ()
        r1 = _rail_row(window, "R1")
        r2 = _rail_row(window, "R2")
        # Workbook Present was stale (1 + 0); current scenario Present wins (3 + 0).
        assert window.distribution_table.item(r1, 1).text() == "3"
        assert window.distribution_table.item(r2, 1).text() == "0"
        assert window.distribution_table.item(r1, 2).text() == "1"
        assert window.distribution_table.item(r2, 2).text() == "2"
        assert window.distribution_distance_combo.currentIndex() == -1
        assert not window.calculate_distribution_button.isEnabled()
        assert "Present refreshed from the loaded SPD: 1 -> 3 (+2" in (
            window.distribution_summary.toPlainText()
        )
        assert "Candidate order was not recorded" in (
            window.distribution_summary.toPlainText()
        )

        window.distribution_distance_combo.setCurrentIndex(0)
        assert window.distribution_distance_combo.currentData() == "NEAREST"
        assert window.calculate_distribution_button.isEnabled()
    finally:
        window._dirty = False
        window.close()
        application.processEvents()


def test_current_target_import_restores_recorded_distance_mode(
    tmp_path: Path,
    monkeypatch,
) -> None:
    application = _application()
    scenario = _direct_scenario(
        (
            ("C1", 0.0, ("R1", "R2")),
            ("C2", 10.0, ("R1", "R2")),
            ("C3", 20.0, ("R1", "R2")),
        ),
        rail_ids=("R1", "R2"),
        bump_x={"R2": 0.0},
    )
    path = tmp_path / "current-targets.xlsx"
    write_distribution_workbook(
        path,
        (),
        (
            "PWR NET",
            "M1\nPresent",
            "M1\nTarget",
            "M1\nTolerance (%)",
            "M1\nActual Delta",
            "M1\nActual Changed",
            "M1\nIsolation Gaps",
        ),
        (
            ("V1 (R1)", 3, 1, 0, -2, 2, 0),
            ("V2 (R2)", 0, 2, 0, 2, 2, 0),
        ),
        metadata={
            "Format Version": 2,
            "Source SPD SHA-256": scenario.source.sha256,
            "Input Design Fingerprint": scenario.design_fingerprint,
            "Distance Mode": "FARTHEST",
        },
    )
    window = _window_with_scenario(scenario)
    try:
        monkeypatch.setattr(
            QFileDialog,
            "getOpenFileName",
            lambda *_args, **_kwargs: (str(path), "Excel workbook (*.xlsx)"),
        )
        window.import_distribution_targets_button.click()

        assert window.distribution_distance_combo.currentData() == "FARTHEST"
        assert window.calculate_distribution_button.isEnabled()
        assert window._distribution_targets[("R1", "M1")] == 1
        assert window._distribution_targets[("R2", "M1")] == 2
    finally:
        window._dirty = False
        window.close()
        application.processEvents()


def test_target_edit_uses_cached_inventory_and_emits_once(monkeypatch) -> None:
    application = _application()
    window = _window_with_scenario(
        _direct_scenario(
            (
                ("C1", 0.0, ("R1", "R2")),
                ("C2", 10.0, ("R1", "R2")),
                ("C3", 20.0, ("R1", "R2")),
            ),
            rail_ids=("R1", "R2"),
        )
    )
    original_getter = ScenarioSpec.base_project.fget
    assert original_getter is not None
    calls = 0

    def counted_base_project(scenario: ScenarioSpec):
        nonlocal calls
        calls += 1
        return original_getter(scenario)

    monkeypatch.setattr(
        ScenarioSpec,
        "base_project",
        property(counted_base_project),
    )
    numeric_calls = 0
    original_balance_state = window._distribution_balance_state

    def counted_balance_state():
        nonlocal numeric_calls
        numeric_calls += 1
        return original_balance_state()

    failure_reset_calls = 0
    original_failure_reset = window._reset_distribution_assignment_failures

    def counted_failure_reset():
        nonlocal failure_reset_calls
        failure_reset_calls += 1
        return original_failure_reset()

    monkeypatch.setattr(
        window,
        "_distribution_balance_state",
        counted_balance_state,
    )
    monkeypatch.setattr(
        window,
        "_reset_distribution_assignment_failures",
        counted_failure_reset,
    )
    spy = QSignalSpy(window.distribution_table.itemChanged)
    try:
        _set_target(window, "R1", 2)
        application.processEvents()

        assert calls == 0
        assert numeric_calls == 1
        assert failure_reset_calls == 0
        assert spy.count() == 1
        assert window._distribution_targets[("R1", "M1")] == 2
    finally:
        window._dirty = False
        window.close()
        application.processEvents()


def test_typing_into_multiselected_targets_fills_all_cells_once() -> None:
    application = _application()
    window = _window_with_scenario(
        _direct_scenario(
            (
                ("C1", 0.0, ("R1", "R2")),
                ("C2", 10.0, ("R1", "R2")),
                ("C3", 20.0, ("R1", "R2")),
            ),
            rail_ids=("R1", "R2"),
        )
    )
    try:
        window.side_tabs.setCurrentIndex(3)
        window.show()
        application.processEvents()

        first = window.distribution_table.item(_rail_row(window, "R1"), 2)
        second = window.distribution_table.item(_rail_row(window, "R2"), 2)
        ignored_present = window.distribution_table.item(_rail_row(window, "R1"), 1)
        ignored_tolerance = window.distribution_table.item(
            _rail_row(window, "R2"), 3
        )
        ignored_delta = window.distribution_table.item(_rail_row(window, "R2"), 4)
        selection = window.distribution_table.selectionModel()
        selection.clearSelection()
        for selected_item in (
            first,
            second,
            ignored_present,
            ignored_tolerance,
            ignored_delta,
        ):
            selection.select(
                window.distribution_table.indexFromItem(selected_item),
                QItemSelectionModel.SelectionFlag.Select,
            )
        window.distribution_table.setCurrentItem(
            first,
            QItemSelectionModel.SelectionFlag.NoUpdate,
        )
        window.distribution_table.setFocus()
        spy = QSignalSpy(window.distribution_table.itemChanged)

        QTest.keyClick(window.distribution_table, Qt.Key.Key_1)
        application.processEvents()
        editor = application.focusWidget()
        assert isinstance(editor, QLineEdit)
        QTest.keyClick(editor, Qt.Key.Key_2)
        QTest.keyClick(editor, Qt.Key.Key_Return)
        application.processEvents()

        assert first.text() == second.text() == "12"
        assert ignored_present.text() == "3"
        assert ignored_tolerance.text() == "0"
        assert ignored_delta.text() == "0"
        assert window._distribution_targets[("R1", "M1")] == 12
        assert window._distribution_targets[("R2", "M1")] == 12
        # The peer-cell fill and styling are signal-blocked, so expensive
        # validation runs for the single committed editor change only.
        assert spy.count() == 1
    finally:
        window._dirty = False
        window.close()
        application.processEvents()


def test_multiselected_fill_runs_when_source_value_is_unchanged() -> None:
    application = _application()
    window = _window_with_scenario(
        _direct_scenario(
            (
                ("C1", 0.0, ("R1", "R2")),
                ("C2", 10.0, ("R1", "R2")),
                ("C3", 20.0, ("R1", "R2")),
            ),
            rail_ids=("R1", "R2"),
        )
    )
    try:
        window.side_tabs.setCurrentIndex(3)
        window.show()
        application.processEvents()

        first = window.distribution_table.item(_rail_row(window, "R1"), 2)
        second = window.distribution_table.item(_rail_row(window, "R2"), 2)
        assert (first.text(), second.text()) == ("3", "0")
        selection = window.distribution_table.selectionModel()
        selection.clearSelection()
        for target in (first, second):
            selection.select(
                window.distribution_table.indexFromItem(target),
                QItemSelectionModel.SelectionFlag.Select,
            )
        window.distribution_table.setCurrentItem(
            first,
            QItemSelectionModel.SelectionFlag.NoUpdate,
        )
        window.distribution_table.setFocus()
        spy = QSignalSpy(window.distribution_table.itemChanged)

        QTest.keyClick(window.distribution_table, Qt.Key.Key_3)
        application.processEvents()
        editor = application.focusWidget()
        assert isinstance(editor, QLineEdit)
        QTest.keyClick(editor, Qt.Key.Key_Return)
        application.processEvents()
        application.processEvents()

        assert first.text() == second.text() == "3"
        assert window._distribution_targets[("R1", "M1")] == 3
        assert window._distribution_targets[("R2", "M1")] == 3
        assert spy.count() == 0
    finally:
        window._dirty = False
        window.close()
        application.processEvents()


def test_multiselected_tolerance_fill_is_decimal_and_never_changes_targets() -> None:
    application = _application()
    window = _window_with_scenario(
        _direct_scenario(
            (
                ("C1", 0.0, ("R1", "R2")),
                ("C2", 10.0, ("R1", "R2")),
            ),
            rail_ids=("R1", "R2"),
        )
    )
    try:
        window.side_tabs.setCurrentIndex(3)
        window.show()
        application.processEvents()

        first = window.distribution_table.item(_rail_row(window, "R1"), 3)
        second = window.distribution_table.item(_rail_row(window, "R2"), 3)
        mixed_target = window.distribution_table.item(_rail_row(window, "R1"), 2)
        original_target = mixed_target.text()
        selection = window.distribution_table.selectionModel()
        selection.clearSelection()
        for selected_item in (first, second, mixed_target):
            selection.select(
                window.distribution_table.indexFromItem(selected_item),
                QItemSelectionModel.SelectionFlag.Select,
            )
        window.distribution_table.setCurrentItem(
            first, QItemSelectionModel.SelectionFlag.NoUpdate
        )
        window.distribution_table.setFocus()
        spy = QSignalSpy(window.distribution_table.itemChanged)

        QTest.keyClick(window.distribution_table, Qt.Key.Key_1)
        application.processEvents()
        editor = application.focusWidget()
        assert isinstance(editor, QLineEdit)
        QTest.keyClicks(editor, ".5")
        QTest.keyClick(editor, Qt.Key.Key_Return)
        application.processEvents()

        assert first.text() == second.text() == "1.5"
        assert window._distribution_tolerances[("R1", "M1")] == 1.5
        assert window._distribution_tolerances[("R2", "M1")] == 1.5
        assert mixed_target.text() == original_target
        assert window._distribution_targets[("R1", "M1")] == int(
            original_target
        )
        assert spy.count() == 1
        assert "no donor/receiver demand" in window.distribution_summary.toPlainText()

        _set_tolerance(window, "R1", "nan")
        assert "finite percentage" in window.distribution_summary.toPlainText()
        assert not window.calculate_distribution_button.isEnabled()
    finally:
        window._dirty = False
        window.close()
        application.processEvents()


def test_numeric_shortage_blocks_calculation_before_physical_planning() -> None:
    application = _application()
    window = _window_with_scenario(
        _direct_scenario(
            (
                ("C1", 0.0, ("R1", "R2")),
                ("C2", 10.0, ("R1", "R2")),
                ("C3", 20.0, ("R1", "R2")),
            ),
            rail_ids=("R1", "R2"),
        )
    )
    try:
        _set_target(window, "R1", 2)  # give capacity 1
        _set_target(window, "R2", 2)  # receive demand 2
        assert not window.calculate_distribution_button.isEnabled()
        assert window.distribution_validation_label.text() == (
            "M1: Donor 1 | Receiver 2 | Balance -1"
        )
        assert "short by 1" in window.distribution_summary.toPlainText()

        _set_target(window, "R1", 1)
        assert window.calculate_distribution_button.isEnabled()
        assert window.distribution_validation_label.text() == (
            "M1: Donor 2 | Receiver 2 | Balance +0"
        )

        _set_target(window, "R2", "1.5")
        assert not window.calculate_distribution_button.isEnabled()
        assert "nonnegative whole number" in window.distribution_summary.toPlainText()
    finally:
        window._dirty = False
        window.close()
        application.processEvents()


def test_distribution_table_sorts_numeric_columns_and_keeps_rail_keys() -> None:
    application = _application()
    scenario = _with_initial_rails(
        _direct_scenario(
            tuple((f"C{index}", float(index), ("R1", "R2")) for index in range(12)),
            rail_ids=("R1", "R2"),
        ),
        {"C0": "R2", "C1": "R2"},
    )
    window = _window_with_scenario(scenario)
    try:
        assert window.distribution_table.isSortingEnabled()
        # Ascending numeric order is 2 then 10; lexical order would be 10 then 2.
        window.distribution_table.sortItems(1, Qt.SortOrder.AscendingOrder)
        assert window.distribution_table.item(0, 0).data(Qt.ItemDataRole.UserRole) == "R2"

        _set_target(window, "R2", 1)
        assert window._distribution_targets[("R2", "M1")] == 1
        assert window._distribution_targets[("R1", "M1")] == 10
    finally:
        window._dirty = False
        window.close()
        application.processEvents()


def test_fixed_donor_capacity_is_not_reported_as_movable_capacity() -> None:
    """The table preflight must match the planner's verified-movable inventory."""

    application = _application()
    scenario = _direct_scenario(
        (
            ("MOVABLE", 0.0, ("R1", "R2")),
            ("FIXED1", 10.0, ("R1", "R2")),
            ("FIXED2", 20.0, ("R1", "R2")),
        ),
        rail_ids=("R1", "R2"),
    )
    assert scenario.connection_analysis is not None
    connections = dict(scenario.connection_analysis.connections)
    for refdes in ("FIXED1", "FIXED2"):
        connections[refdes] = ScenarioDecapConnection(
            refdes=refdes,
            kind=DecapConnectionKind.UNRESOLVED,
            reason="source topology is ambiguous",
        )
    scenario = ScenarioSpec.model_validate(
        {
            **scenario.model_dump(mode="python"),
            "connection_analysis": scenario.connection_analysis.model_copy(
                update={"connections": connections}
            ),
        }
    )
    window = _window_with_scenario(scenario)
    try:
        # Physical inventory says this donor can give two. Only MOVABLE has
        # verified connectivity, so numeric donor capacity is one.
        _set_target(window, "R1", 1)
        _set_target(window, "R2", 2)

        assert not window.calculate_distribution_button.isEnabled()
        assert window.distribution_validation_label.text() == (
            "M1: Donor 1 | Receiver 2 | Balance -1"
        )
        message = window.distribution_summary.toPlainText()
        assert "donor 1" in message
        assert "receiver 2" in message
        assert "fixed/unassignable 1" in message
        assert "short by 1" in message
    finally:
        window._dirty = False
        window.close()
        application.processEvents()


def test_gnd_only_unresolved_donor_enables_calculation_with_exact_proof_pending() -> None:
    application = _application()
    source = _shared_chain_scenario()
    assert source.connection_analysis is not None
    connections = {
        refdes: connection.model_copy(
            update={
                "kind": DecapConnectionKind.UNRESOLVED,
                "reason": (
                    "GND TOP component has no source Via anchor: "
                    f"{refdes}"
                ),
            }
        )
        for refdes, connection in source.connection_analysis.connections.items()
    }
    cluster = source.connection_analysis.clusters[0].model_copy(
        update={
            "state": SharedPadClusterState.UNRESOLVED,
            "ground_edges": (),
            "isolation_gap_refdes": (),
            "reason": "GND TOP component has no source Via anchor: A0, D1, A2",
            "eligibility": {},
            "via_eligibility": {},
        }
    )
    decaps = tuple(
        item.model_copy(
            update={
                "pwr_pad": item.pwr_pad.model_copy(update={"padstack": "PAD"}),
                "gnd_pad": item.gnd_pad.model_copy(update={"padstack": "PAD"}),
            }
        )
        for item in source.decaps
    )
    scenario = ScenarioSpec.model_validate(
        {
            **source.model_dump(mode="python"),
            "decaps": decaps,
            "connection_analysis": source.connection_analysis.model_copy(
                update={"connections": connections, "clusters": (cluster,)}
            ),
        }
    )
    window = _window_with_scenario(scenario)
    try:
        _set_target(window, "R1", 0)
        _set_target(window, "R2", 3)

        assert window.calculate_distribution_button.isEnabled()
        assert window.distribution_validation_label.text() == (
            "M1: Donor 3 | Receiver 3 | Balance +0"
        )
        message = window.distribution_summary.toPlainText()
        assert "exact retained-artwork proof pending" in message
        assert "donor 3" in message
        assert "receiver 3" in message
    finally:
        window._dirty = False
        window.close()
        application.processEvents()


def test_distribution_worker_reuses_one_projection_for_validate_compute_and_apply(
    monkeypatch,
) -> None:
    import spd_decap_pi.distribution as distribution_module

    scenario = _direct_scenario(
        (("C1", 0.0, ("R1", "R2")),),
        rail_ids=("R1", "R2"),
    )
    projection = object()
    plan = SimpleNamespace()
    seen: list[tuple[str, object | None]] = []

    def build(*_args, **kwargs):
        assert kwargs["targets"] == {("R1", "M1"): 0, ("R2", "M1"): 1}
        assert kwargs["tolerances"] == {
            ("R1", "M1"): 0.0,
            ("R2", "M1"): 0.0,
        }
        seen.append(("build", None))
        return projection

    def validate(*_args, **kwargs):
        seen.append(("validate", kwargs.get("power_projection")))

    def compute(*_args, **kwargs):
        seen.append(("compute", kwargs.get("power_projection")))
        return plan

    def apply(*_args, **kwargs):
        seen.append(("apply", kwargs.get("power_projection")))
        return scenario

    monkeypatch.setattr(
        distribution_module, "build_distribution_power_projection", build
    )
    monkeypatch.setattr(distribution_module, "validate_distribution_targets", validate)
    monkeypatch.setattr(distribution_module, "compute_distribution_plan", compute)
    monkeypatch.setattr(distribution_module, "apply_distribution_plan", apply)
    monkeypatch.setattr(
        distribution_module,
        "distribution_csv_rows",
        lambda _plan: (("Component", "REFDES"), ("M1", "C1")),
    )

    result = _job_compute_distribution(
        scenario,
        {"geometry/test": b"payload"},
        {("R1", "M1"): 0, ("R2", "M1"): 1},
        {("R1", "M1"): 0.0, ("R2", "M1"): 0.0},
        DistributionDistanceMode.NEAREST,
        progress=lambda _value, _message: None,
        is_cancelled=lambda: False,
    )

    assert seen == [
        ("build", None),
        ("validate", projection),
        ("compute", projection),
        ("apply", projection),
    ]
    assert result.power_projection is projection
    assert result.preview_scenario is scenario


def test_candidate_audit_uses_receiver_gross_inbound_counterflow_demand(
    monkeypatch,
) -> None:
    import spd_decap_pi.distribution as distribution_module

    scenario = _with_initial_rails(
        _direct_scenario(
            (
                ("A0", 0.0, ("R1", "R2")),
                ("A1", 10.0, ("R1", "R2")),
                ("A2", 20.0, ("R1", "R2")),
                ("B0", 100.0, ("R2", "R3")),
            ),
            bump_x={"R2": 0.0, "R3": 100.0},
        ),
        {"B0": "R2"},
    )
    monkeypatch.setattr(
        distribution_module,
        "build_distribution_power_projection",
        lambda *_args, **_kwargs: None,
    )

    prepared = _job_compute_distribution(
        scenario,
        {},
        {("R1", "M1"): 1, ("R2", "M1"): 2, ("R3", "M1"): 1},
        {("R1", "M1"): 0.0, ("R2", "M1"): 100.0, ("R3", "M1"): 0.0},
        DistributionDistanceMode.NEAREST,
        progress=lambda _value, _message: None,
        is_cancelled=lambda: False,
    )

    receiver = next(
        cell for cell in prepared.plan.cells if cell.rail_id == "R2"
    )
    assert (receiver.requested_count, receiver.sent_count, receiver.received_count) == (
        1,
        1,
        2,
    )
    unselected = next(
        row
        for row in prepared.candidate_audit_rows
        if row[0] == "A2" and row[4] == "R2"
    )
    assert unselected[9] == "ELIGIBLE_NOT_SELECTED_BY_GLOBAL_OPTIMUM"
    assert "(1x2=2)" in unselected[10]


def test_candidate_audit_keeps_partial_receiver_full_demand(
    monkeypatch,
) -> None:
    import spd_decap_pi.distribution_audit as audit_module

    scenario = _direct_scenario(
        (
            ("C1", 0.0, ("R1", "R2")),
            ("C2", 10.0, ("R1",)),
            ("C3", 20.0, ("R1",)),
        ),
        rail_ids=("R1", "R2"),
        bump_x={"R2": 0.0},
    )
    original = audit_module.audit_distribution_candidates
    calls: list[tuple[str, int]] = []

    def record_demand(*args, **kwargs):
        calls.append((str(args[1]), int(args[2])))
        return original(*args, **kwargs)

    monkeypatch.setattr(audit_module, "audit_distribution_candidates", record_demand)
    prepared = _job_compute_distribution(
        scenario,
        {},
        {("R1", "M1"): 1, ("R2", "M1"): 2},
        {("R1", "M1"): 0.0, ("R2", "M1"): 0.0},
        DistributionDistanceMode.NEAREST,
        progress=lambda _value, _message: None,
        is_cancelled=lambda: False,
    )

    receiver = next(
        cell for cell in prepared.plan.cells if cell.rail_id == "R2"
    )
    assert (receiver.requested_count, receiver.received_count) == (2, 0)
    assert ("R2", 2) in calls
    assert any(row[4] == "R2" for row in prepared.candidate_audit_rows)


def test_one_invalid_component_blocks_an_otherwise_valid_component() -> None:
    application = _application()
    base = _direct_scenario(
        tuple(
            (f"C{index}", float(index), ("R1", "R2"))
            for index in range(1, 7)
        ),
        rail_ids=("R1", "R2"),
    )
    project = base.base_project
    m2 = project.cap_models[0].model_copy(
        update={"model_id": "M2", "source_hash": "m2"}
    )
    scenario = ScenarioSpec.model_validate(
        {
            **base.model_dump(mode="python"),
            "normalized_project": project.model_copy(
                update={"cap_models": [*project.cap_models, m2]}
            ),
            "decaps": [
                item
                if index < 3
                else item.model_copy(
                    update={"model_id": "M2", "source_model_id": "M2"}
                )
                for index, item in enumerate(base.decaps)
            ],
        }
    )
    window = _window_with_scenario(scenario)
    try:
        # M1: capacity 2, demand 2 -> valid.
        _set_target(window, "R1", 1, model_id="M1")
        _set_target(window, "R2", 2, model_id="M1")
        # M2: capacity 1, demand 2 -> invalid and blocks the whole calculation.
        _set_target(window, "R1", 2, model_id="M2")
        _set_target(window, "R2", 2, model_id="M2")

        assert not window.calculate_distribution_button.isEnabled()
        assert "M2 is short by 1" in window.distribution_summary.toPlainText()
        assert "M1: Donor 2 | Receiver 2 | Balance +0" in (
            window.distribution_validation_label.text()
        )
    finally:
        window._dirty = False
        window.close()
        application.processEvents()


def test_full_preview_exports_saves_and_applies_one_atomic_revision(
    tmp_path: Path,
    monkeypatch,
) -> None:
    application = _application()
    scenario = _direct_scenario(
        (
            ("C1", 0.0, ("R1", "R2")),
            ("C2", 10.0, ("R1", "R2")),
            ("C3", 20.0, ("R1", "R2")),
        ),
        rail_ids=("R1", "R2"),
        bump_x={"R2": 0.0},
    )
    window = _window_with_scenario(scenario)
    try:
        _set_target(window, "R1", 1)
        _set_target(window, "R2", 2)
        plan = compute_distribution_plan(
            scenario,
            _targets(window),
            DistributionDistanceMode.NEAREST,
        )
        window._accept_distribution_plan(plan)

        assert plan.status == DistributionPlanStatus.FULL
        assert window.apply_distribution_button.isEnabled()
        assert window.export_distribution_csv_button.isEnabled()
        assert window.save_distribution_button.isEnabled()
        assert window.distribution_table.item(_rail_row(window, "R1"), 4).text() == "0"
        assert window.distribution_table.item(_rail_row(window, "R2"), 4).text() == "0"
        summary = window.distribution_summary.toPlainText()
        assert "Status: FULL (count/topology)" in summary
        assert (
            "PDN evaluation (modified/touched rails): no inherited "
            "unresolved/out-of-scope connection blockers"
            in summary
        )
        assert "R1 donor: give capacity 2, used 2, unused 0" in summary
        assert "R2 receiver: requested 2, fulfilled 2, shortfall 0" in summary

        export_path = tmp_path / "distribution"
        monkeypatch.setattr(
            QFileDialog,
            "getSaveFileName",
            lambda *_args, **_kwargs: (str(export_path), "CSV files (*.csv)"),
        )
        window.export_distribution_csv_button.click()
        actual_csv = export_path.with_suffix(".csv")
        assert actual_csv.read_bytes().startswith(b"\xef\xbb\xbf")
        with actual_csv.open("r", encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream))
        assert rows[0] == [
            "Component",
            "REFDES",
            "Before NET",
            "After NET",
            "X (um)",
            "Y (um)",
        ]
        assert len(rows) == len(scenario.decaps) + 1
        assert {row[1] for row in rows[1:]} == {"C1", "C2", "C3"}

        excel_path = tmp_path / "distribution.xlsx"
        monkeypatch.setattr(
            QFileDialog,
            "getSaveFileName",
            lambda *_args, **_kwargs: (
                str(excel_path),
                "Excel workbook (*.xlsx)",
            ),
        )
        window.export_distribution_csv_button.click()
        workbook = load_workbook(excel_path, data_only=False)
        try:
            assert workbook.sheetnames == [
                "Decap Changes",
                "PWR NET Distribution Targets",
            ]
            targets = workbook["PWR NET Distribution Targets"]
            assert tuple(cell.value for cell in targets[1][:11]) == (
                "PWR NET",
                "M1\nPresent",
                "M1\nTarget",
                "M1\nTolerance (%)",
                "M1\nRole",
                "M1\nTurnover Allowance",
                "M1\nSent",
                "M1\nReceived",
                "M1\nActual Delta",
                "M1\nAssignment Failed",
                "M1\nIsolation Gaps",
            )
            assert tuple(cell.value for cell in targets[2][:11]) == (
                "V1 (R1)",
                3,
                1,
                0,
                "DONOR",
                0,
                2,
                0,
                -2,
                0,
                0,
            )
            assert tuple(cell.value for cell in targets[3][:11]) == (
                "V2 (R2)",
                0,
                2,
                0,
                "RECEIVER",
                0,
                0,
                2,
                2,
                0,
                0,
            )
            metadata_title_row = next(
                cell.row
                for cell in targets["A"]
                if cell.value == DISTRIBUTION_METADATA_TITLE
            )
            metadata = {}
            for row_index in range(metadata_title_row + 1, targets.max_row + 1):
                key = targets.cell(row_index, 1).value
                if key in (None, ""):
                    break
                metadata[str(key)] = targets.cell(row_index, 2).value
            assert metadata["Format Version"] == 5
            assert metadata["Signal Routing Protection"] == "OFF"
            assert (
                metadata["Tolerance Semantics"]
                == DISTRIBUTION_TOLERANCE_SEMANTICS
            )
            assert metadata["Source SPD SHA-256"] == scenario.source.sha256
            assert metadata["Input Design Fingerprint"] == scenario.design_fingerprint
            assert metadata["Distance Mode"] == "NEAREST"
            assert metadata["Optimization Policy"] == "BALANCED_AUTO"
            assert metadata["Via Projection Policy"] == DISTRIBUTION_VIA_PROJECTION_POLICY
            assert metadata["Effective Gap Penalty (um)"] == pytest.approx(
                plan.effective_gap_penalty_um
            )
        finally:
            workbook.close()

        def run_immediately(worker, on_result, **_kwargs):
            result = worker.function(
                *worker.args,
                progress=lambda _value, _message: None,
                is_cancelled=lambda: False,
                **worker.kwargs,
            )
            on_result(result)

        distributed_path = tmp_path / "preview"
        monkeypatch.setattr(window, "_run_worker", run_immediately)
        monkeypatch.setattr(
            QFileDialog,
            "getSaveFileName",
            lambda *_args, **_kwargs: (
                str(distributed_path),
                "SPD PI Scenario (*.spdpi)",
            ),
        )
        window.save_distribution_button.click()
        persisted = load_scenario(distributed_path.with_suffix(".spdpi"))
        assert persisted.revision == scenario.revision + 1
        assert sum(item.current_rail_id == "R2" for item in persisted.decaps) == 2
        assert window.scenario is scenario

        window.apply_distribution_button.click()
        assert window.scenario is not None
        assert window.scenario.revision == scenario.revision + 1
        assert sum(item.current_rail_id == "R2" for item in window.scenario.decaps) == 2
        assert window._dirty
        assert not window.apply_distribution_button.isEnabled()
        assert window.export_distribution_csv_button.isEnabled()
        assert window.save_distribution_button.isEnabled()
        # The applied table shows final Present values, preserves the requested
        # Target, and retains the preview's assignment-failure result.
        assert window.distribution_table.item(_rail_row(window, "R1"), 1).text() == "1"
        assert window.distribution_table.item(_rail_row(window, "R1"), 2).text() == "1"
        assert window.distribution_table.item(_rail_row(window, "R1"), 4).text() == "0"
        assert window.distribution_table.item(_rail_row(window, "R2"), 1).text() == "2"
        assert window.distribution_table.item(_rail_row(window, "R2"), 2).text() == "2"
        assert window.distribution_table.item(_rail_row(window, "R2"), 4).text() == "0"

        applied_excel_path = tmp_path / "distribution-applied.xlsx"
        monkeypatch.setattr(
            QFileDialog,
            "getSaveFileName",
            lambda *_args, **_kwargs: (
                str(applied_excel_path),
                "Excel workbook (*.xlsx)",
            ),
        )
        window.export_distribution_csv_button.click()
        applied_workbook = load_workbook(applied_excel_path, data_only=False)
        try:
            targets = applied_workbook["PWR NET Distribution Targets"]
            # Apply refreshes the GUI inventory to final Present values, but
            # the exported target sheet remains the immutable plan input.
            assert tuple(cell.value for cell in targets[2][:11]) == (
                "V1 (R1)",
                3,
                1,
                0,
                "DONOR",
                0,
                2,
                0,
                -2,
                0,
                0,
            )
            assert tuple(cell.value for cell in targets[3][:11]) == (
                "V2 (R2)",
                0,
                2,
                0,
                "RECEIVER",
                0,
                0,
                2,
                2,
                0,
                0,
            )
        finally:
            applied_workbook.close()

        applied_path = tmp_path / "applied-distribution"
        monkeypatch.setattr(
            QFileDialog,
            "getSaveFileName",
            lambda *_args, **_kwargs: (
                str(applied_path),
                "SPD PI Scenario (*.spdpi)",
            ),
        )
        window.save_distribution_button.click()
        assert window._scenario_path == applied_path.with_suffix(".spdpi")
        assert not window._dirty
        assert load_scenario(window._scenario_path).design_fingerprint == (
            window.scenario.design_fingerprint
        )
    finally:
        window._dirty = False
        window.close()
        application.processEvents()


def test_full_preview_separates_topology_success_from_evaluation_block() -> None:
    application = _application()
    scenario = _direct_scenario(
        (
            ("C1", 0.0, ("R1", "R2")),
            ("C2", 10.0, ("R1", "R2")),
        ),
        rail_ids=("R1", "R2"),
        bump_x={"R2": 0.0},
    )
    window = _window_with_scenario(scenario)
    try:
        _set_target(window, "R1", 1)
        _set_target(window, "R2", 1)
        plan = compute_distribution_plan(
            scenario,
            _targets(window),
            DistributionDistanceMode.NEAREST,
        )
        plan = replace(
            plan,
            diagnostics=plan.diagnostics
            + (
                DistributionDiagnostic(
                    code="PREEXISTING_UNRESOLVED_EVALUATION_RAILS",
                    message=(
                        "count/topology preview is valid, but inherited "
                        "connection evidence blocks PDN evaluation"
                    ),
                ),
            ),
        )

        window._accept_distribution_plan(plan)

        summary = window.distribution_summary.toPlainText()
        assert "Status: FULL (count/topology)" in summary
        assert (
            "PDN evaluation (modified/touched rails): BLOCKED on inherited "
            "unresolved/out-of-scope connections"
            in summary
        )
        assert "PDN evaluation blocked" in window.status_text.text()
        assert window.apply_distribution_button.isEnabled()
    finally:
        window._dirty = False
        window.close()
        application.processEvents()


def test_run_evaluation_preflights_all_selected_connectivity_blockers_without_worker(
    monkeypatch,
) -> None:
    application = _application()
    scenario = _with_initial_rails(
        _direct_scenario(
            (
                ("C1", 0.0, ("R1",)),
                ("C2", 10.0, ("R2",)),
            ),
            rail_ids=("R1", "R2"),
        ),
        {"C2": "R2"},
    )
    payload = scenario.model_dump(mode="python")
    for refdes, kind, reason in (
        ("C1", DecapConnectionKind.UNRESOLVED, "R1 source landing is unresolved"),
        ("C2", DecapConnectionKind.OUT_OF_SCOPE, "R2 component is outside TOP scope"),
    ):
        payload["connection_analysis"]["connections"][refdes].update(
            {
                "kind": kind,
                "power_vias": (),
                "ground_vias": (),
                "reason": reason,
            }
        )
    blocked = ScenarioSpec.model_validate(payload)
    window = _window_with_scenario(blocked)
    warnings: list[tuple[str, str, str]] = []
    workers: list[tuple[object, object]] = []
    try:
        for index in range(window.rail_list.count()):
            window.rail_list.item(index).setCheckState(Qt.CheckState.Checked)
        monkeypatch.setattr(
            QMessageBox,
            "exec",
            lambda box: warnings.append(
                (box.text(), box.informativeText(), box.detailedText())
            ),
        )
        monkeypatch.setattr(
            window,
            "_run_worker",
            lambda worker, on_result, **_kwargs: workers.append(
                (worker, on_result)
            ),
        )
        window._comparison_batch = object()
        window.comparison_table.setRowCount(1)

        window.run_evaluation()

        assert len(workers) == 1
        preflight_worker, accept_preflight = workers[0]
        assert preflight_worker.function.__name__ == "_job_preflight_evaluation"
        accept_preflight(
            preflight_evaluation_connectivity(blocked, ("R1", "R2"))
        )
        assert len(workers) == 1  # all blocked: no Evaluation Analysis worker
        assert len(warnings) == 1
        summary, guidance, details = warnings[0]
        assert summary == (
            "Evaluation cannot start: 2 decap connectivity/modelability blocker(s) "
            "affect 2 of 2 selected PWR rail(s)."
        )
        assert "R1" not in summary and "C1" not in summary
        assert "Show Details" in guidance and "Evaluation Summary" in guidance
        assert "R1" in details and "C1" in details
        assert "R2" in details and "C2" in details
        assert window._comparison_batch is None
        assert window.comparison_table.rowCount() == 0
        evaluation_summary = window.evaluation_summary.toPlainText()
        assert "blocked before Original/Tuned" in evaluation_summary
        assert details in evaluation_summary
    finally:
        window._dirty = False
        window.close()
        application.processEvents()


def test_run_evaluation_preflights_disabled_and_isolation_gap_blockers_without_worker(
    monkeypatch,
) -> None:
    application = _application()
    scenario = _direct_scenario((("C1", 0.0, ("R1",)), ("C2", 10.0, ("R1",))), rail_ids=("R1",))
    gap = scenario.decaps[0].model_copy(
        update={"enabled": False, "pad_state": DecapPadState.ISOLATION_GAP}
    )
    disabled = scenario.decaps[1].model_copy(update={"enabled": False})
    connections = dict(scenario.connection_analysis.connections)
    connections["C1"] = connections["C1"].model_copy(
        update={
            "kind": DecapConnectionKind.UNRESOLVED,
            "power_vias": (),
            "ground_vias": (),
            "reason": "isolation gap source landing is unresolved",
        }
    )
    connections["C2"] = connections["C2"].model_copy(
        update={
            "kind": DecapConnectionKind.OUT_OF_SCOPE,
            "power_vias": (),
            "ground_vias": (),
            "reason": "disabled pad is outside TOP scope",
        }
    )
    # Direct isolation gaps are not persistable, but model_copy keeps this
    # focused GUI regression on the builder's connection-before-pad-state path.
    blocked = scenario.model_copy(
        update={
            "decaps": [gap, disabled],
            "connection_analysis": scenario.connection_analysis.model_copy(
                update={"connections": connections}
            ),
        }
    )
    window = _window_with_scenario(blocked)
    warnings: list[tuple[str, str]] = []
    workers: list[tuple[object, object]] = []
    try:
        monkeypatch.setattr(
            QMessageBox,
            "exec",
            lambda box: warnings.append((box.text(), box.detailedText())),
        )
        monkeypatch.setattr(
            window,
            "_run_worker",
            lambda worker, on_result, **_kwargs: workers.append(
                (worker, on_result)
            ),
        )

        window.run_evaluation()

        assert len(workers) == 1
        preflight_worker, accept_preflight = workers[0]
        assert preflight_worker.function.__name__ == "_job_preflight_evaluation"
        accept_preflight(preflight_evaluation_connectivity(blocked, ("R1",)))
        assert len(workers) == 1  # all blocked: no Evaluation Analysis worker
        assert len(warnings) == 1
        summary, details = warnings[0]
        assert "1 of 1 selected PWR rail(s)" in summary
        assert "C1" not in summary and "C2" not in summary
        assert "C1" in details and "C2" in details
        assert "UNRESOLVED" in details and "OUT_OF_SCOPE" in details
    finally:
        window._dirty = False
        window.close()
        application.processEvents()


def test_partial_evaluation_requires_yes_and_worker_receives_only_clear_bare_rail(
    tmp_path: Path, monkeypatch,
) -> None:
    """A VQPS-style bare rail may run, but a blocked peer is never hidden."""

    application = _application()
    scenario = _direct_scenario(
        (("C1", 0.0, ("R1",)),),
        rail_ids=("R1", "R2"),
    )
    payload = scenario.model_dump(mode="python")
    payload["connection_analysis"]["connections"]["C1"].update(
        {
            "kind": DecapConnectionKind.UNRESOLVED,
            "power_vias": (),
            "ground_vias": (),
            "reason": "R1 source landing is unresolved",
        }
    )
    blocked = ScenarioSpec.model_validate(payload)
    window = _window_with_scenario(blocked)
    window._scenario_path = tmp_path / "partial.spdpi"
    workers: list[tuple[object, object]] = []
    prompts: list[tuple[str, str, str]] = []
    try:
        window._set_all_rails_checked(True)
        monkeypatch.setattr(
            QMessageBox,
            "exec",
            lambda box: (
                prompts.append(
                    (box.text(), box.informativeText(), box.detailedText())
                )
                or QMessageBox.StandardButton.Yes
            ),
        )
        monkeypatch.setattr(
            window,
            "_run_worker",
            lambda worker, on_result, **_kwargs: workers.append(
                (worker, on_result)
            ),
        )

        window.run_evaluation()
        assert len(workers) == 1
        workers[0][1](
            preflight_evaluation_connectivity(blocked, ("R1", "R2"))
        )

        assert len(prompts) == 1
        title, guidance, details = prompts[0]
        assert title == "Run a partial Evaluation Analysis?"
        assert "1 clear PWR NET(s) will run" in guidance
        assert "1 blocked PWR NET(s) will NOT run" in guidance
        assert "R2" in details and "Clear and evaluated" in details
        assert "R1" in details and "Blocked and NOT evaluated" in details

        request, manifest = window._pending_evaluation_launch
        window._pending_evaluation_launch = None
        window._launch_evaluation_after_preflight(request, manifest)
        assert len(workers) == 2
        evaluation_worker = workers[1][0]
        assert evaluation_worker.function.__name__ == "evaluate_comparison_batch"
        assert evaluation_worker.args[1] == ("R2",)
        assert set(evaluation_worker.args[0].baseline_captures) == {"R2"}
        summary = window.evaluation_summary.toPlainText()
        assert "Selected PWR NETs: 2" in summary
        assert "Blocked and NOT evaluated: 1 (R1)" in summary
    finally:
        window._dirty = False
        window.close()
        application.processEvents()


def test_partial_evaluation_no_cancels_without_analysis_worker(
    tmp_path: Path, monkeypatch,
) -> None:
    application = _application()
    scenario = _direct_scenario(
        (("C1", 0.0, ("R1",)),),
        rail_ids=("R1", "R2"),
    )
    payload = scenario.model_dump(mode="python")
    payload["connection_analysis"]["connections"]["C1"].update(
        {
            "kind": DecapConnectionKind.UNRESOLVED,
            "power_vias": (),
            "ground_vias": (),
            "reason": "R1 source landing is unresolved",
        }
    )
    blocked = ScenarioSpec.model_validate(payload)
    window = _window_with_scenario(blocked)
    window._scenario_path = tmp_path / "partial-no.spdpi"
    workers: list[tuple[object, object]] = []
    try:
        window._set_all_rails_checked(True)
        monkeypatch.setattr(
            QMessageBox,
            "exec",
            lambda _box: QMessageBox.StandardButton.No,
        )
        monkeypatch.setattr(
            window,
            "_run_worker",
            lambda worker, on_result, **_kwargs: workers.append(
                (worker, on_result)
            ),
        )

        window.run_evaluation()
        workers[0][1](
            preflight_evaluation_connectivity(blocked, ("R1", "R2"))
        )

        assert len(workers) == 1
        assert window._pending_evaluation_launch is None
        assert "cancelled by user" in window.evaluation_summary.toPlainText()
        assert "1 clear / 1 blocked" in window.status_text.text()
    finally:
        window._dirty = False
        window.close()
        application.processEvents()


def test_partial_preview_remains_applyable_and_target_edits_make_it_stale() -> None:
    application = _application()
    scenario = _direct_scenario(
        (
            ("C1", 0.0, ("R1", "R2")),
            ("C2", 10.0, ("R1",)),
            ("C3", 20.0, ("R1",)),
        ),
        rail_ids=("R1", "R2"),
        bump_x={"R2": 0.0},
    )
    window = _window_with_scenario(scenario)
    try:
        _set_target(window, "R1", 1)
        _set_target(window, "R2", 2)
        plan = compute_distribution_plan(scenario, _targets(window))
        window._accept_distribution_plan(plan)

        assert plan.status == DistributionPlanStatus.PARTIAL
        assert "Status: PARTIAL" in window.distribution_summary.toPlainText()
        assert "shortfall 1" in window.distribution_summary.toPlainText()
        assert "R1 donor: give capacity 2, used 1, unused 1" in (
            window.distribution_summary.toPlainText()
        )
        assert "R2 receiver: requested 2, fulfilled 1, shortfall 1" in (
            window.distribution_summary.toPlainText()
        )
        assert window.apply_distribution_button.isEnabled()
        assert window.distribution_table.item(_rail_row(window, "R1"), 4).text() == "0"
        assert window.distribution_table.item(_rail_row(window, "R2"), 4).text() == "1"

        window.apply_distribution_button.click()
        assert window.scenario is not None
        assert window.scenario.revision == scenario.revision + 1
        assert sum(item.current_rail_id == "R2" for item in window.scenario.decaps) == 1
        assert window.distribution_table.item(_rail_row(window, "R2"), 2).text() == "2"
        assert window.distribution_table.item(_rail_row(window, "R2"), 4).text() == "1"

        _set_target(window, "R2", 3)
        assert window._distribution_plan is None
        assert not window.apply_distribution_button.isEnabled()
        assert not window.export_distribution_csv_button.isEnabled()
        assert not window.save_distribution_button.isEnabled()
    finally:
        window._dirty = False
        window.close()
        application.processEvents()


def test_exchange_preview_reports_turnover_and_apply_preserves_tolerance() -> None:
    application = _application()
    scenario = _with_initial_rails(
        _direct_scenario(
            (
                ("A0", 0.0, ("R1", "R2")),
                ("A1", 10.0, ("R1", "R2")),
                ("B0", 100.0, ("R2", "R3")),
                ("B1", 110.0, ("R2", "R3")),
            ),
            bump_x={"R2": 100.0, "R3": 200.0},
        ),
        {"B0": "R2", "B1": "R2"},
    )
    window = _window_with_scenario(scenario)
    try:
        _set_target(window, "R1", 1)
        _set_target(window, "R2", 2)
        _set_target(window, "R3", 1)
        _set_tolerance(window, "R2", 50)
        assert window.calculate_distribution_button.isEnabled()
        assert "tolerance counterflow 1 cell(s) / 1 decap(s) allowance" in (
            window.distribution_summary.toPlainText()
        )

        plan = compute_distribution_plan(
            scenario,
            _targets(window),
            tolerances=dict(window._distribution_tolerances),
        )
        window._accept_distribution_plan(plan)

        assert plan.status == DistributionPlanStatus.FULL
        assert plan.changed_count == 2
        assert window.distribution_table.item(_rail_row(window, "R2"), 4).text() == "0"
        assert (
            "M1 / R2 exchange: tolerance 50% (1), sent 1, received 1, net +0"
            in window.distribution_summary.toPlainText()
        )

        window.apply_distribution_button.click()
        assert window._distribution_tolerances[("R2", "M1")] == 50.0
        assert window.distribution_table.item(_rail_row(window, "R2"), 3).text() == "50"
    finally:
        window._dirty = False
        window.close()
        application.processEvents()


def test_distribution_milp_worker_allows_cancel_at_safe_solver_boundaries(
    monkeypatch,
) -> None:
    application = _application()
    window = _window_with_scenario(
        _direct_scenario(
            (
                ("C1", 0.0, ("R1", "R2")),
                ("C2", 10.0, ("R1", "R2")),
                ("C3", 20.0, ("R1", "R2")),
            ),
            rail_ids=("R1", "R2"),
            bump_x={"R2": 0.0},
        )
    )
    captured: dict[str, object] = {}
    try:
        _set_target(window, "R1", 1)
        _set_target(window, "R2", 2)
        _set_tolerance(window, "R2", 1.25)

        def capture_worker(worker, _on_result, **kwargs):
            captured["worker"] = worker
            captured.update(kwargs)

        monkeypatch.setattr(window, "_run_worker", capture_worker)
        window.calculate_distribution_button.click()

        assert captured["label"] == "Calculating De-cap Distribution preview..."
        assert captured["cancelable"] is True
        worker = captured["worker"]
        assert worker.args[3][("R2", "M1")] == 1.25
    finally:
        window._dirty = False
        window.close()
        application.processEvents()


def test_completed_distribution_worker_does_not_leave_calculating_status() -> None:
    application = _application()
    window = MainWindow()
    try:
        window.status_text.setText("Calculating De-cap Distribution preview...")

        window._worker_finished()

        assert window.status_text.text() == "Ready"
    finally:
        window.close()
        application.processEvents()


def test_scenario_change_and_document_reset_discard_a_stale_preview() -> None:
    application = _application()
    scenario = _direct_scenario(
        (
            ("C1", 0.0, ("R1", "R2")),
            ("C2", 10.0, ("R1", "R2")),
            ("C3", 20.0, ("R1", "R2")),
        ),
        rail_ids=("R1", "R2"),
        bump_x={"R2": 0.0},
    )
    window = _window_with_scenario(scenario)
    try:
        _set_target(window, "R1", 1)
        _set_target(window, "R2", 2)
        window._accept_distribution_plan(
            compute_distribution_plan(scenario, _targets(window))
        )
        assert window._distribution_plan is not None

        changed = ScenarioSpec.model_validate(
            {
                **scenario.model_dump(mode="python"),
                "decaps": [
                    scenario.decaps[0].model_copy(update={"enabled": False}),
                    *scenario.decaps[1:],
                ],
                "revision": scenario.revision + 1,
            }
        )
        window._scenario = changed
        window._refresh_all()
        assert window._distribution_plan is None
        assert window._distribution_preview_scenario is None
        assert not window.export_distribution_csv_button.isEnabled()
        assert "every Target equals Present" in window.distribution_summary.toPlainText()

        window._reset_document_view_state()
        assert window.distribution_table.rowCount() == 0
        assert window._distribution_basis_fingerprint is None
    finally:
        window._dirty = False
        window.close()
        application.processEvents()
