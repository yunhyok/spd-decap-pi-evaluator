from __future__ import annotations

from pathlib import Path
import re
import zipfile

from openpyxl import load_workbook
import pytest

import spd_decap_pi.spreadsheet_export as spreadsheet_export
from spd_decap_pi.distribution_workbook import load_distribution_targets
from spd_decap_pi.spreadsheet_export import (
    CANDIDATE_AUDIT_HEADERS,
    EXCEL_MAX_DATA_ROWS,
    write_distribution_workbook,
)


def _column_widths(path: Path, sheet_part: str) -> dict[int, float]:
    """Return 1-based column index -> emitted width for one worksheet part.

    ``openpyxl`` does not expand ``<col min=.. max=../>`` ranges, so the raw
    worksheet XML is the only faithful view of the exported widths.
    """

    with zipfile.ZipFile(path) as archive:
        xml = archive.read(sheet_part).decode("utf-8")
    widths: dict[int, float] = {}
    for element in re.findall(r"<col [^>]*/>", xml):
        first = int(re.search(r'min="(\d+)"', element).group(1))
        last = int(re.search(r'max="(\d+)"', element).group(1))
        width = float(re.search(r'width="([\d.]+)"', element).group(1))
        for column in range(first, last + 1):
            widths[column] = width
    return widths


def test_distribution_workbook_has_two_typed_formula_safe_sheets(
    tmp_path: Path,
) -> None:
    path = tmp_path / "distribution.xlsx"
    write_distribution_workbook(
        path,
        (
            ("=MODEL", "+C1", "@BEFORE", "-AFTER", 1.25, -2.5),
            ("M1", "C2", "V1", "V1", 3.0, 4.0),
        ),
        (
            "PWR NET",
            "M1\nPresent",
            "M1\nTarget",
            "M1\nTolerance (%)",
            "M1\nActual Delta",
        ),
        (
            ("V1 (R1)", 2, 1, 1.25, -1),
            ("V2 (R2)", 0, 1, 0.0, 1),
        ),
    )

    workbook = load_workbook(path, data_only=False)
    try:
        assert workbook.sheetnames == [
            "Decap Changes",
            "PWR NET Distribution Targets",
        ]
        decaps = workbook["Decap Changes"]
        assert tuple(cell.value for cell in decaps[1]) == (
            "Component",
            "REFDES",
            "Before NET",
            "After NET",
            "X (um)",
            "Y (um)",
        )
        assert tuple(cell.value for cell in decaps[2]) == (
            "=MODEL",
            "+C1",
            "@BEFORE",
            "-AFTER",
            1.25,
            -2.5,
        )
        assert all(decaps.cell(2, column).data_type == "s" for column in range(1, 5))
        assert all(decaps.cell(2, column).data_type == "n" for column in range(5, 7))
        assert decaps.freeze_panes == "A2"
        assert decaps.auto_filter.ref == "A1:F3"

        targets = workbook["PWR NET Distribution Targets"]
        assert tuple(cell.value for cell in targets[1]) == (
            "PWR NET",
            "M1\nPresent",
            "M1\nTarget",
            "M1\nTolerance (%)",
            "M1\nActual Delta",
        )
        assert tuple(cell.value for cell in targets[2]) == (
            "V1 (R1)",
            2,
            1,
            1.25,
            -1,
        )
        assert tuple(cell.value for cell in targets[3]) == (
            "V2 (R2)",
            0,
            1,
            0,
            1,
        )
        assert targets.cell(2, 4).number_format == '0.###"%"'
        assert targets.freeze_panes == "B2"
        assert targets.auto_filter.ref == "A1:E3"
    finally:
        workbook.close()


def test_distribution_workbook_rejects_misaligned_target_rows(
    tmp_path: Path,
) -> None:
    path = tmp_path / "invalid.xlsx"
    try:
        write_distribution_workbook(path, (), ("PWR NET", "Target"), (("R1",),))
    except ValueError as exc:
        assert "match their headers" in str(exc)
    else:
        raise AssertionError("misaligned target row was accepted")
    assert not path.exists()


def test_distribution_workbook_optionally_writes_candidate_audit_sheet(
    tmp_path: Path,
) -> None:
    path = tmp_path / "distribution-with-audit.xlsx"
    write_distribution_workbook(
        path,
        (),
        ("PWR NET", "M1\nTarget"),
        (("REF_CLK (R1)", 15),),
        candidate_audit_headers=CANDIDATE_AUDIT_HEADERS,
        candidate_audit_rows=(
            (
                "C793-C805",
                "C793, C794",
                13,
                "SRC",
                "REF_CLK",
                "M1",
                12.5,
                True,
                False,
                "NO_ZERO_GAP_EXACT_COUNT_COMBINATION",
                "component size 13 cannot be part of any zero-gap exact-count subset",
            ),
        ),
    )
    workbook = load_workbook(path, data_only=False)
    try:
        assert workbook.sheetnames == [
            "Decap Changes",
            "PWR NET Distribution Targets",
            "Candidate Audit",
        ]
        audit = workbook["Candidate Audit"]
        assert audit.cell(1, 1).value == "RefDes"
        assert audit.cell(2, 1).value == "C793-C805"
        assert audit.cell(2, 10).value == "NO_ZERO_GAP_EXACT_COUNT_COMBINATION"
        assert audit.freeze_panes == "A2"
    finally:
        workbook.close()
    imported = load_distribution_targets(
        path,
        rail_ids=("R1",),
        model_ids=("M1",),
        current_present={("R1", "M1"): 4},
    )
    assert imported.targets == {("R1", "M1"): 15}


def test_distribution_workbook_formats_distribution_columns_for_multiple_models(
    tmp_path: Path,
) -> None:
    path = tmp_path / "distribution-five-column-groups.xlsx"
    headers = ["PWR NET"]
    row: list[object] = ["V1 (R1)"]
    for model_id, tolerance, delta, gaps in (
        ("M1", 1.25, -2, 1),
        ("M2", 2.5, 3, 0),
    ):
        headers.extend(
            (
                f"{model_id}\nPresent",
                f"{model_id}\nTarget",
                f"{model_id}\nTolerance (%)",
                f"{model_id}\nActual Delta",
                f"{model_id}\nAssignment Failed",
                f"{model_id}\nIsolation Gaps",
            )
        )
        row.extend((10, 8, tolerance, delta, 4, gaps))
    write_distribution_workbook(path, (), tuple(headers), (tuple(row),))

    workbook = load_workbook(path, data_only=False)
    try:
        targets = workbook["PWR NET Distribution Targets"]
        assert targets.cell(2, 4).number_format == '0.###"%"'
        assert targets.cell(2, 5).number_format == "+#,##0;-#,##0;0"
        assert targets.cell(2, 6).number_format == "#,##0"
        assert targets.cell(2, 7).number_format == "#,##0"
        assert targets.cell(2, 10).number_format == '0.###"%"'
        assert targets.cell(2, 11).number_format == "+#,##0;-#,##0;0"
        assert targets.cell(2, 12).number_format == "#,##0"
        assert targets.cell(2, 13).number_format == "#,##0"
    finally:
        workbook.close()


def test_candidate_audit_row_limit_is_checked_before_xlsxwriter_truncation(
    tmp_path: Path,
) -> None:
    row = ("C1", "C1", 1, "SRC", "DST", "M1", 1.0, True, False, "X", "detail")
    path = tmp_path / "overflow.xlsx"
    with pytest.raises(ValueError, match="worksheet limit"):
        write_distribution_workbook(
            path,
            (),
            ("PWR NET",),
            (("R1",),),
            candidate_audit_headers=CANDIDATE_AUDIT_HEADERS,
            candidate_audit_rows=(row,) * (EXCEL_MAX_DATA_ROWS + 1),
        )
    assert not path.exists()


def test_candidate_audit_row_limit_boundary_is_accepted(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # Avoid writing a million physical rows in a unit test; validation occurs
    # before population and this verifies the exact Excel boundary is allowed.
    monkeypatch.setattr(spreadsheet_export, "_populate_distribution_workbook", lambda *args: None)
    path = tmp_path / "boundary.xlsx"
    row = ("C1", "C1", 1, "SRC", "DST", "M1", 1.0, True, False, "X", "detail")
    write_distribution_workbook(
        path,
        (),
        ("PWR NET",),
        (("R1",),),
        candidate_audit_headers=CANDIDATE_AUDIT_HEADERS,
        candidate_audit_rows=(row,) * EXCEL_MAX_DATA_ROWS,
    )
    assert path.exists()


def test_writer_refuses_the_policy_penalty_pairings_the_loader_rejects(
    tmp_path: Path,
) -> None:
    headers = ("PWR NET", "M1\nTarget")
    rows = (("V1 (R1)", 3),)

    missing = tmp_path / "balanced-custom-without-penalty.xlsx"
    with pytest.raises(ValueError, match="Effective Gap Penalty"):
        write_distribution_workbook(
            missing,
            (),
            headers,
            rows,
            metadata={"Optimization Policy": "BALANCED_CUSTOM"},
        )
    assert not missing.exists()

    stray = tmp_path / "min-gaps-with-penalty.xlsx"
    with pytest.raises(ValueError, match="must be 0 or omitted"):
        write_distribution_workbook(
            stray,
            (),
            headers,
            rows,
            metadata={
                "Optimization Policy": "MIN_GAPS",
                "Effective Gap Penalty (um)": 120.0,
            },
        )
    assert not stray.exists()

    accepted = tmp_path / "balanced-custom.xlsx"
    write_distribution_workbook(
        accepted,
        (),
        headers,
        rows,
        metadata={
            "Optimization Policy": "BALANCED_CUSTOM",
            "Effective Gap Penalty (um)": 120.0,
        },
    )
    imported = load_distribution_targets(
        accepted,
        rail_ids=("R1",),
        model_ids=("M1",),
        current_present={("R1", "M1"): 1},
    )
    assert imported.targets == {("R1", "M1"): 3}


def test_failed_export_leaves_the_previous_workbook_intact(tmp_path: Path) -> None:
    path = tmp_path / "distribution.xlsx"
    write_distribution_workbook(
        path,
        (("M1", "C1", "V1", "V2", 1.0, 2.0),),
        ("PWR NET", "M1\nTarget"),
        (("V1 (R1)", 3),),
    )
    original = path.read_bytes()

    with pytest.raises(ValueError, match="must be finite"):
        write_distribution_workbook(
            path,
            (("M1", "C1", "V1", "V2", float("nan"), 2.0),),
            ("PWR NET", "M1\nTarget"),
            (("V1 (R1)", 4),),
        )

    assert path.read_bytes() == original
    assert sorted(item.name for item in tmp_path.iterdir()) == [path.name]
    imported = load_distribution_targets(
        path,
        rail_ids=("R1",),
        model_ids=("M1",),
        current_present={("R1", "M1"): 1},
    )
    assert imported.targets == {("R1", "M1"): 3}


def test_metadata_value_width_does_not_leak_into_the_target_matrix(
    tmp_path: Path,
) -> None:
    path = tmp_path / "metadata-widths.xlsx"
    headers = ["PWR NET"]
    row: list[object] = ["V1 (R1)"]
    for model_id in ("M1", "M2", "M3"):
        headers.extend(
            (
                f"{model_id}\nPresent",
                f"{model_id}\nTarget",
                f"{model_id}\nTolerance (%)",
            )
        )
        row.extend((10, 8, 1.25))
    write_distribution_workbook(
        path,
        (),
        tuple(headers),
        (tuple(row),),
        metadata={
            "Optimization Policy": "BALANCED_AUTO",
            "Distance Mode": "NEAREST",
        },
    )

    widths = _column_widths(path, "xl/worksheets/sheet2.xml")
    data_columns = range(2, len(headers) + 1)
    assert {widths[column] for column in data_columns} == {widths[2]}
    assert widths[1] != widths[2]


def test_metadata_value_width_does_not_leak_with_inventory_rows(
    tmp_path: Path,
) -> None:
    path = tmp_path / "inventory-widths.xlsx"
    write_distribution_workbook(
        path,
        (),
        ("PWR NET", "M1\nPresent", "M1\nTarget"),
        (("V1 (R1)", 2, 3),),
        inventory_headers=("Component", "Physical Present", "Delta"),
        inventory_rows=(("M1", 2, 1),),
        metadata={"Optimization Policy": "BALANCED_AUTO"},
    )

    widths = _column_widths(path, "xl/worksheets/sheet2.xml")
    assert widths[2] == widths[3] == 16.7109375
