from __future__ import annotations

import math

import numpy as np
from openpyxl import load_workbook
import pytest

import spd_decap_pi.distribution as distribution_module
from spd_decap_pi.distribution_workbook import load_distribution_targets
from spd_decap_pi.spreadsheet_export import write_distribution_workbook

from spd_decap_pi.distribution import (
    DistributionOptimizationPolicy,
    compute_distribution_plan,
)
from spd_decap_pi.scenario import (
    DecapConnectionKind,
    ScenarioDecapConnection,
    ScenarioSpec,
)

from test_spd_decap_distribution import (
    _decap,
    _shared_chain_scenario,
    _via,
)


_NEAR_GAP_DISTANCE_UM = 71_976.014
_FAR_ZERO_GAP_DISTANCE_UM = 281_164.628
_REAL_DISTANCE_SAVING_UM = (
    _FAR_ZERO_GAP_DISTANCE_UM - _NEAR_GAP_DISTANCE_UM
)


def _gap_distance_tradeoff_scenario(
    *,
    near_distance_um: float = _NEAR_GAP_DISTANCE_UM,
    far_distance_um: float = _FAR_ZERO_GAP_DISTANCE_UM,
) -> ScenarioSpec:
    base = _shared_chain_scenario()
    relocated = []
    for offset, decap in enumerate(base.decaps):
        x_um = near_distance_um + float(offset)
        relocated.append(
            decap.model_copy(
                update={
                    "center": decap.center.model_copy(update={"x_um": x_um}),
                    "pwr_pad": decap.pwr_pad.model_copy(
                        update={"x_um": x_um - 10.0}
                    ),
                    "gnd_pad": decap.gnd_pad.model_copy(
                        update={"x_um": x_um + 10.0}
                    ),
                }
            )
        )
    far = _decap("FAR", far_distance_um, ("R1", "R2"))
    assert base.connection_analysis is not None
    connections = dict(base.connection_analysis.connections)
    connections[far.refdes] = ScenarioDecapConnection(
        refdes=far.refdes,
        kind=DecapConnectionKind.DIRECT,
        power_vias=(_via("VP-FAR", "V1", far_distance_um),),
        ground_vias=(_via("VG-FAR", "DGND", far_distance_um + 5.0),),
    )
    project = base.base_project.model_copy(
        update={
            "outline": base.base_project.outline.model_copy(
                update={"width_um": 99_400.0, "height_um": 99_400.0}
            )
        }
    )
    analysis = base.connection_analysis.model_copy(
        update={"connections": connections}
    )
    return ScenarioSpec.model_validate(
        {
            **base.model_dump(mode="python"),
            "normalized_project": project.model_dump(mode="python"),
            "decaps": [
                *(item.model_dump(mode="python") for item in relocated),
                far.model_dump(mode="python"),
            ],
            "connection_analysis": analysis.model_dump(mode="python"),
        }
    )


def _targets() -> dict[tuple[str, str], int]:
    return {("R1", "M1"): 1, ("R2", "M1"): 2}


def test_balanced_auto_records_board_diagonal_penalty() -> None:
    scenario = _shared_chain_scenario()
    plan = compute_distribution_plan(scenario, _targets())

    assert plan.optimization_policy == DistributionOptimizationPolicy.BALANCED_AUTO
    assert plan.effective_gap_penalty_um == pytest.approx(
        math.hypot(scenario.base_project.outline.width_um, scenario.base_project.outline.height_um)
    )


def test_balanced_custom_and_min_gaps_are_explicit() -> None:
    scenario = _shared_chain_scenario()
    custom = compute_distribution_plan(
        scenario,
        _targets(),
        optimization_policy=DistributionOptimizationPolicy.BALANCED_CUSTOM,
        gap_penalty_um=1234.5,
    )
    legacy = compute_distribution_plan(
        scenario,
        _targets(),
        optimization_policy=DistributionOptimizationPolicy.MIN_GAPS,
    )

    assert custom.optimization_policy == DistributionOptimizationPolicy.BALANCED_CUSTOM
    assert custom.effective_gap_penalty_um == pytest.approx(1234.5)
    assert legacy.optimization_policy == DistributionOptimizationPolicy.MIN_GAPS
    assert legacy.effective_gap_penalty_um == 0.0


def test_balanced_auto_accepts_one_gap_when_distance_saving_exceeds_diagonal() -> None:
    scenario = _gap_distance_tradeoff_scenario()
    plan = compute_distribution_plan(
        scenario,
        {("R1", "M1"): 1, ("R2", "M1"): 1},
    )

    assert plan.effective_gap_penalty_um == pytest.approx(
        math.hypot(99_400.0, 99_400.0)
    )
    assert plan.effective_gap_penalty_um < _REAL_DISTANCE_SAVING_UM
    assert plan.moves[0].refdes != "FAR"
    assert len(plan.isolation_gap_refdes) == 1


def test_custom_penalty_above_distance_saving_selects_zero_gap_candidate() -> None:
    scenario = _gap_distance_tradeoff_scenario()
    plan = compute_distribution_plan(
        scenario,
        {("R1", "M1"): 1, ("R2", "M1"): 1},
        optimization_policy="BALANCED_CUSTOM",
        gap_penalty_um=_REAL_DISTANCE_SAVING_UM + 1.0,
    )

    assert plan.moves[0].refdes == "FAR"
    assert plan.isolation_gap_refdes == ()


def test_balanced_auto_prefers_zero_gap_when_distances_are_close() -> None:
    scenario = _gap_distance_tradeoff_scenario(
        near_distance_um=100_000.0,
        far_distance_um=110_000.0,
    )
    plan = compute_distribution_plan(
        scenario,
        {("R1", "M1"): 1, ("R2", "M1"): 1},
    )

    assert plan.moves[0].refdes == "FAR"
    assert plan.isolation_gap_refdes == ()


def test_min_gaps_and_farthest_keep_signed_policy_behavior() -> None:
    scenario = _gap_distance_tradeoff_scenario()
    legacy = compute_distribution_plan(
        scenario,
        {("R1", "M1"): 1, ("R2", "M1"): 1},
        optimization_policy="MIN_GAPS",
    )
    farthest = compute_distribution_plan(
        scenario,
        {("R1", "M1"): 1, ("R2", "M1"): 1},
        distance_mode="FARTHEST",
    )

    assert legacy.moves[0].refdes == "FAR"
    assert legacy.isolation_gap_refdes == ()
    assert farthest.moves[0].refdes == "FAR"
    assert farthest.isolation_gap_refdes == ()


def test_one_milli_um_combined_advantage_survives_gap_and_canonical_ties() -> None:
    scenario = _gap_distance_tradeoff_scenario(
        near_distance_um=1_000.0,
        far_distance_um=1_100.001,
    )
    plan = compute_distribution_plan(
        scenario,
        {("R1", "M1"): 1, ("R2", "M1"): 1},
        optimization_policy="BALANCED_CUSTOM",
        gap_penalty_um=100.0,
    )

    # gap1 is better by exactly one integer milli-um objective unit.  Later
    # gap-count and canonical solves are constrained to that proven level and
    # therefore cannot reverse it in favour of FAR/gap0.
    assert plan.moves[0].refdes != "FAR"
    assert len(plan.isolation_gap_refdes) == 1


def test_balanced_solver_keeps_distance_coefficients_at_physical_scale(
    monkeypatch,
) -> None:
    observed_maxima: list[float] = []
    original = distribution_module._milp_with_optional_start

    def capture_objective(c, **kwargs):
        values = np.abs(np.asarray(c, dtype=float))
        observed_maxima.append(float(np.max(values)) if values.size else 0.0)
        return original(c, **kwargs)

    monkeypatch.setattr(
        distribution_module,
        "_milp_with_optional_start",
        capture_objective,
    )
    compute_distribution_plan(
        _gap_distance_tradeoff_scenario(),
        {("R1", "M1"): 1, ("R2", "M1"): 1},
    )

    assert observed_maxima
    assert max(observed_maxima) <= round(_FAR_ZERO_GAP_DISTANCE_UM * 1000.0)


def test_workbook_policy_and_penalty_round_trip(tmp_path) -> None:
    path = tmp_path / "policy.xlsx"
    write_distribution_workbook(
        path,
        (),
        ("PWR NET", "M1\nTarget"),
        (("V1 (R1)", 1),),
        metadata={
            "Format Version": 3,
            "Optimization Policy": "BALANCED_CUSTOM",
            "Effective Gap Penalty (um)": 1234.5,
        },
    )
    imported = load_distribution_targets(
        path,
        rail_ids=("R1",),
        model_ids=("M1",),
        current_present={("R1", "M1"): 1},
    )
    assert imported.optimization_policy == "BALANCED_CUSTOM"
    assert imported.effective_gap_penalty_um == pytest.approx(1234.5)


@pytest.mark.parametrize(
    ("policy", "penalty", "message"),
    (
        ("BALANCED_CUSTOM", None, "missing Effective Gap Penalty"),
        ("BALANCED_CUSTOM", 1.0e308, "from 0 through"),
        ("MIN_GAPS", 1.0, "must be 0 or omitted"),
    ),
)
def test_workbook_rejects_incomplete_or_contradictory_policy_metadata(
    tmp_path,
    policy: str,
    penalty: float | None,
    message: str,
) -> None:
    path = tmp_path / f"invalid-{policy}.xlsx"
    # ``write_distribution_workbook`` refuses these policy/penalty pairings
    # itself, so the invalid workbook has to be built by writing a valid one
    # and then rewriting the two metadata cells in place.  That keeps this
    # test on the loader gate it targets instead of the writer gate.
    write_distribution_workbook(
        path,
        (),
        ("PWR NET", "M1\nTarget"),
        (("V1 (R1)", 1),),
        metadata={
            "Format Version": 3,
            "Optimization Policy": "BALANCED_CUSTOM",
            "Effective Gap Penalty (um)": 1.0,
        },
    )
    workbook = load_workbook(path, data_only=False)
    sheet = workbook["PWR NET Distribution Targets"]
    for row in sheet.iter_rows(min_col=1, max_col=2):
        if row[0].value == "Optimization Policy":
            row[1].value = policy
        elif row[0].value == "Effective Gap Penalty (um)":
            row[1].value = penalty
    workbook.save(path)

    with pytest.raises(ValueError, match=message):
        load_distribution_targets(
            path,
            rail_ids=("R1",),
            model_ids=("M1",),
            current_present={("R1", "M1"): 1},
        )


@pytest.mark.parametrize(
    "penalty", (-1.0, float("nan"), float("inf"), 1.0e308)
)
def test_balanced_custom_rejects_invalid_penalty(penalty: float) -> None:
    with pytest.raises(ValueError, match="gap penalty"):
        compute_distribution_plan(
            _shared_chain_scenario(),
            _targets(),
            optimization_policy="BALANCED_CUSTOM",
            gap_penalty_um=penalty,
        )
