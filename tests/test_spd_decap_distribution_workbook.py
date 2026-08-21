from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

import spd_decap_pi.distribution_workbook as distribution_workbook
from spd_decap_pi.distribution_workbook import (
    DISTRIBUTION_METADATA_TITLE,
    DISTRIBUTION_TOLERANCE_SEMANTICS,
    DISTRIBUTION_VIA_PROJECTION_POLICY,
    DistributionWorkbookError,
    load_distribution_targets,
)
from spd_decap_pi.spreadsheet_export import write_distribution_workbook


LEGACY_HEADERS = (
    "PWR NET",
    "M1\nPresent",
    "M1\nTarget",
    "M1\nTolerance (%)",
    "M1\nActual Delta",
)

CURRENT_OFF_METADATA = {
    "Format Version": 5,
    "Signal Routing Protection": "OFF",
    "Tolerance Semantics": DISTRIBUTION_TOLERANCE_SEMANTICS,
    "Via Projection Policy": DISTRIBUTION_VIA_PROJECTION_POLICY,
}


def _current_metadata(**updates: object) -> dict[str, object]:
    result = dict(CURRENT_OFF_METADATA)
    result.update(updates)
    return result


def test_legacy_targets_are_absolute_and_present_is_refreshed_from_current_scenario(
    tmp_path: Path,
) -> None:
    path = tmp_path / "legacy.xlsx"
    write_distribution_workbook(
        path,
        (),
        LEGACY_HEADERS,
        (
            ("V1 (R1)", 2, 1, 0.0, "=ignored-result"),
            ("V2 (R2)", 1, 2, 0.0, 1),
        ),
        metadata=_current_metadata(),
    )

    imported = load_distribution_targets(
        path,
        rail_ids=("R1", "R2", "R3"),
        model_ids=("M1",),
        current_present={
            ("R1", "M1"): 4,
            ("R2", "M1"): 1,
            ("R3", "M1"): 2,
        },
    )

    assert imported.targets == {
        ("R1", "M1"): 1,
        ("R2", "M1"): 2,
        ("R3", "M1"): 2,
    }
    assert imported.tolerances == {
        ("R1", "M1"): 0.0,
        ("R2", "M1"): 0.0,
        ("R3", "M1"): 0.0,
    }
    assert imported.workbook_present_total == 3
    assert imported.current_present_total == 7
    assert imported.changed_present_cells == 1
    assert imported.defaulted_current_cells == 1
    assert imported.distance_mode is None
    assert "source identity was not recorded" in imported.summary(path.name)
    assert "3 -> 7 (+4" in imported.summary(path.name)


def test_current_export_keeps_two_sheets_a1_matrix_and_round_trips_metadata(
    tmp_path: Path,
) -> None:
    path = tmp_path / "current.xlsx"
    source_sha = "a" * 64
    fingerprint = "b" * 64
    headers = (
        "PWR NET",
        "M1\nPresent",
        "M1\nTarget",
        "M1\nTolerance (%)",
        "M1\nRole",
        "M1\nTurnover Allowance",
        "M1\nSent",
        "M1\nReceived",
        "M1\nActual Delta",
        "M1\nAssignment Failed",
        "M1\nIsolation Gaps",
    )
    write_distribution_workbook(
        path,
        (("M1", "C1", "V1", "V2", 1.0, 2.0),),
        headers,
        (("V1 (R1)", 4, 3, 1.25, "DONOR", 0, 1, 0, -1, 2, 1),),
        inventory_headers=("Component", "Physical Present"),
        inventory_rows=(("M1", 4),),
        metadata={
            "Format Version": 5,
            "Signal Routing Protection": "OFF",
            "Tolerance Semantics": DISTRIBUTION_TOLERANCE_SEMANTICS,
            "Application Version": "0.9.3",
            "Source SPD SHA-256": source_sha,
            "Input Design Fingerprint": fingerprint,
            "Distance Mode": "FARTHEST",
            "Via Projection Policy": DISTRIBUTION_VIA_PROJECTION_POLICY,
        },
    )

    workbook = load_workbook(path, data_only=False)
    try:
        assert workbook.sheetnames == [
            "Decap Changes",
            "PWR NET Distribution Targets",
        ]
        sheet = workbook["PWR NET Distribution Targets"]
        assert tuple(cell.value for cell in sheet[1][: len(headers)]) == headers
        assert sheet["A1"].value == "PWR NET"
        assert any(
            cell.value == DISTRIBUTION_METADATA_TITLE for cell in sheet["A"]
        )
    finally:
        workbook.close()

    imported = load_distribution_targets(
        path,
        rail_ids=("R1",),
        model_ids=("M1",),
        current_present={("R1", "M1"): 5},
        current_source_sha256=source_sha,
        current_design_fingerprint=fingerprint,
    )
    assert imported.targets[("R1", "M1")] == 3
    assert imported.tolerances[("R1", "M1")] == 1.25
    assert imported.distance_mode == "FARTHEST"
    assert imported.format_version == 5
    assert imported.source_sha256 == source_sha
    assert imported.via_projection_policy == DISTRIBUTION_VIA_PROJECTION_POLICY


def test_legacy_format_requires_reexport_with_source_proven_policy(tmp_path: Path) -> None:
    path = tmp_path / "legacy-no-source.xlsx"
    write_distribution_workbook(
        path,
        (),
        ("PWR NET", "M1\nTarget"),
        (("V1 (R1)", 1),),
        metadata={"Format Version": 3},
    )

    with pytest.raises(DistributionWorkbookError, match="requires re-export"):
        load_distribution_targets(
            path,
            rail_ids=("R1",),
            model_ids=("M1",),
            current_present={("R1", "M1"): 1},
        )


def test_unknown_via_projection_policy_fails_closed(tmp_path: Path) -> None:
    path = tmp_path / "unknown-via-projection.xlsx"
    write_distribution_workbook(
        path,
        (),
        ("PWR NET", "M1\nTarget"),
        (("V1 (R1)", 1),),
        metadata={
            "Format Version": 5,
            "Tolerance Semantics": DISTRIBUTION_TOLERANCE_SEMANTICS,
            "Signal Routing Protection": "OFF",
            "Via Projection Policy": "REQUIRE_EXISTING_MLO_PATH",
        },
    )

    with pytest.raises(DistributionWorkbookError, match="Via Projection Policy"):
        load_distribution_targets(
            path,
            rail_ids=("R1",),
            model_ids=("M1",),
            current_present={("R1", "M1"): 1},
        )


def test_writer_rejects_format4_without_routing_mode_metadata(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="Signal Routing Protection"):
        write_distribution_workbook(
            tmp_path / "invalid-v4.xlsx",
            (),
            ("PWR NET", "M1\nTarget"),
            (("V1 (R1)", 1),),
            metadata={"Format Version": 4},
        )


def test_format5_requires_explicit_tolerance_semantics(tmp_path: Path) -> None:
    path = tmp_path / "missing-tolerance-semantics.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "PWR NET Distribution Targets"
    sheet.append(("PWR NET", "M1\nPresent", "M1\nTarget", "M1\nTolerance (%)"))
    sheet.append(("V1 (R1)", 3, 2, 10.0))
    sheet.append(())
    sheet.append((DISTRIBUTION_METADATA_TITLE,))
    sheet.append(("Format Version", 5))
    sheet.append(("Signal Routing Protection", "OFF"))
    workbook.save(path)
    workbook.close()

    with pytest.raises(DistributionWorkbookError, match="Tolerance Semantics"):
        load_distribution_targets(
            path,
            rail_ids=("R1",),
            model_ids=("M1",),
            current_present={("R1", "M1"): 3},
        )


def test_current_target_changing_tolerance_round_trips(
    tmp_path: Path,
) -> None:
    path = tmp_path / "legacy-directional-tolerance.xlsx"
    write_distribution_workbook(
        path,
        (),
        ("PWR NET", "M1\nPresent", "M1\nTarget", "M1\nTolerance (%)"),
        (("V1 (R1)", 3, 2, 10.0),),
        metadata=_current_metadata(),
    )

    imported = load_distribution_targets(
        path,
        rail_ids=("R1",),
        model_ids=("M1",),
        current_present={("R1", "M1"): 3},
    )
    assert imported.tolerances[("R1", "M1")] == 10.0


def test_legacy_safe_tolerance_cells_still_load(tmp_path: Path) -> None:
    path = tmp_path / "legacy-safe-tolerance.xlsx"
    write_distribution_workbook(
        path,
        (),
        ("PWR NET", "M1\nPresent", "M1\nTarget", "M1\nTolerance (%)"),
        (
            ("V1 (R1)", 3, 3, 10.0),
            ("V2 (R2)", 2, 1, 0.0),
        ),
        metadata=_current_metadata(),
    )

    imported = load_distribution_targets(
        path,
        rail_ids=("R1", "R2"),
        model_ids=("M1",),
        current_present={("R1", "M1"): 3, ("R2", "M1"): 2},
    )

    assert imported.tolerances[("R1", "M1")] == 10.0
    assert imported.targets[("R2", "M1")] == 1


def test_legacy_migration_uses_current_present_not_stale_workbook_present(
    tmp_path: Path,
) -> None:
    path = tmp_path / "legacy-stale-present.xlsx"
    write_distribution_workbook(
        path,
        (),
        ("PWR NET", "M1\nPresent", "M1\nTarget", "M1\nTolerance (%)"),
        (("V1 (R1)", 3, 3, 10.0),),
        metadata=_current_metadata(),
    )

    imported = load_distribution_targets(
        path,
        rail_ids=("R1",),
        model_ids=("M1",),
        current_present={("R1", "M1"): 4},
    )
    assert imported.targets[("R1", "M1")] == 3


@pytest.mark.parametrize("result_header", ("Actual Changed", "Assignment Failed"))
def test_import_ignores_legacy_and_current_assignment_result_columns(
    tmp_path: Path,
    result_header: str,
) -> None:
    path = tmp_path / f"result-{result_header.casefold().replace(' ', '-')}.xlsx"
    headers = (
        "PWR NET",
        "M1\nPresent",
        "M1\nTarget",
        f"M1\n{result_header}",
    )
    write_distribution_workbook(
        path,
        (),
        headers,
        (("V1 (R1)", 3, 2, 99),),
        metadata=_current_metadata(),
    )

    imported = load_distribution_targets(
        path,
        rail_ids=("R1",),
        model_ids=("M1",),
        current_present={("R1", "M1"): 3},
    )

    assert imported.targets == {("R1", "M1"): 2}


def test_import_streams_bounded_ranges_and_skips_result_columns_for_target_rows(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "bounded.xlsx"
    write_distribution_workbook(
        path,
        (("M1", "C1", "V1", "V2", 1.0, 2.0),),
        LEGACY_HEADERS,
        (("V1 (R1)", 3, 2, 0, -1),),
        inventory_headers=("Component", "Physical Present"),
        inventory_rows=(("M1", 3),),
        metadata=_current_metadata(**{"Distance Mode": "NEAREST"}),
    )

    calls: list[dict[str, object]] = []
    original_iter_rows = ReadOnlyWorksheet.iter_rows

    def bounded_iter_rows(self, *args, **kwargs):
        calls.append(dict(kwargs))
        return original_iter_rows(self, *args, **kwargs)

    monkeypatch.setattr(ReadOnlyWorksheet, "iter_rows", bounded_iter_rows)
    imported = load_distribution_targets(
        path,
        rail_ids=("R1",),
        model_ids=("M1",),
        current_present={("R1", "M1"): 3},
    )

    assert imported.targets[("R1", "M1")] == 2
    assert imported.distance_mode == "NEAREST"
    assert len(calls) == 3
    assert all(call.get("max_row") is not None for call in calls)
    assert all(call.get("max_col") is not None for call in calls)
    assert all(
        int(call["max_row"]) <= distribution_workbook._MAX_IMPORT_ROWS
        for call in calls
    )
    assert all(
        int(call["max_col"]) <= distribution_workbook._MAX_IMPORT_COLUMNS
        for call in calls
    )
    assert calls[0]["min_row"] == calls[0]["max_row"] == 1
    assert calls[1]["min_row"] == 2
    assert calls[1]["max_col"] == 4
    assert calls[2]["max_col"] == 2


@pytest.mark.parametrize(
    ("far_row", "far_column"),
    (
        (distribution_workbook._MAX_IMPORT_ROWS + 1, 1),
        (1, distribution_workbook._MAX_IMPORT_COLUMNS + 1),
    ),
)
def test_unreasonable_sheet_dimensions_fail_before_cell_iteration(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    far_row: int,
    far_column: int,
) -> None:
    path = tmp_path / f"oversized-{far_row}-{far_column}.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "PWR NET Distribution Targets"
    sheet.append(LEGACY_HEADERS)
    sheet.append(("V1 (R1)", 3, 2, 0, -1))
    sheet.cell(row=far_row, column=far_column, value="stray used-range cell")
    workbook.save(path)
    workbook.close()

    def unexpected_iter_rows(*_args, **_kwargs):
        raise AssertionError("dimension guard must run before worksheet iteration")

    monkeypatch.setattr(ReadOnlyWorksheet, "iter_rows", unexpected_iter_rows)
    with pytest.raises(DistributionWorkbookError, match="dimensions are unreasonable"):
        load_distribution_targets(
            path,
            rail_ids=("R1",),
            model_ids=("M1",),
            current_present={("R1", "M1"): 3},
        )


def test_unmatched_active_target_fails_closed_but_neutral_cell_is_ignored(
    tmp_path: Path,
) -> None:
    active_path = tmp_path / "active.xlsx"
    write_distribution_workbook(
        active_path,
        (),
        LEGACY_HEADERS,
        (("OLD (REMOVED)", 2, 1, 0, -1),),
        metadata=_current_metadata(),
    )
    with pytest.raises(DistributionWorkbookError, match="active workbook target"):
        load_distribution_targets(
            active_path,
            rail_ids=("R1",),
            model_ids=("M1",),
            current_present={("R1", "M1"): 3},
        )

    neutral_path = tmp_path / "neutral.xlsx"
    write_distribution_workbook(
        neutral_path,
        (),
        LEGACY_HEADERS,
        (
            ("V1 (R1)", 3, 2, 0, -1),
            ("OLD (REMOVED)", 4, 4, 0, 0),
        ),
        metadata=_current_metadata(),
    )
    imported = load_distribution_targets(
        neutral_path,
        rail_ids=("R1",),
        model_ids=("M1",),
        current_present={("R1", "M1"): 3},
    )
    assert imported.targets[("R1", "M1")] == 2
    assert imported.ignored_neutral_cells == 1


def test_duplicate_rail_rows_fail_closed(tmp_path: Path) -> None:
    path = tmp_path / "duplicate-rail.xlsx"
    write_distribution_workbook(
        path,
        (),
        LEGACY_HEADERS,
        (
            ("V1 (R1)", 3, 2, 0, -1),
            ("V1 duplicate (r1)", 3, 2, 0, -1),
        ),
    )
    with pytest.raises(DistributionWorkbookError, match="duplicate rail row"):
        load_distribution_targets(
            path,
            rail_ids=("R1",),
            model_ids=("M1",),
            current_present={("R1", "M1"): 3},
        )


@pytest.mark.parametrize(
    ("mutator", "message"),
    (
        (lambda sheet: sheet.cell(1, 6, "M1\nTarget"), "duplicate target"),
        (lambda sheet: sheet.cell(2, 3, -1), "nonnegative whole"),
        (lambda sheet: sheet.cell(2, 4, 101), "0 through 100"),
        (lambda sheet: sheet.cell(2, 3, "=1+1"), "formulas are not supported"),
    ),
)
def test_duplicate_and_invalid_instruction_cells_fail_closed(
    tmp_path: Path,
    mutator,
    message: str,
) -> None:
    path = tmp_path / "invalid.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "PWR NET Distribution Targets"
    sheet.append(LEGACY_HEADERS)
    sheet.append(("V1 (R1)", 3, 2, 0, -1))
    mutator(sheet)
    workbook.save(path)
    workbook.close()

    with pytest.raises(DistributionWorkbookError, match=message):
        load_distribution_targets(
            path,
            rail_ids=("R1",),
            model_ids=("M1",),
            current_present={("R1", "M1"): 3},
        )


def test_metadata_source_mismatch_and_unknown_distance_fail_closed(
    tmp_path: Path,
) -> None:
    path = tmp_path / "metadata-invalid.xlsx"
    write_distribution_workbook(
        path,
        (),
        LEGACY_HEADERS,
        (("V1 (R1)", 3, 2, 0, -1),),
        metadata={
            **_current_metadata(),
            "Source SPD SHA-256": "a" * 64,
            "Distance Mode": "SIDEWAYS",
        },
    )
    with pytest.raises(DistributionWorkbookError, match="Distance Mode"):
        load_distribution_targets(
            path,
            rail_ids=("R1",),
            model_ids=("M1",),
            current_present={("R1", "M1"): 3},
            current_source_sha256="b" * 64,
        )

    path2 = tmp_path / "source-invalid.xlsx"
    write_distribution_workbook(
        path2,
        (),
        LEGACY_HEADERS,
        (("V1 (R1)", 3, 2, 0, -1),),
        metadata={
            **_current_metadata(),
            "Source SPD SHA-256": "a" * 64,
            "Distance Mode": "NEAREST",
        },
    )
    with pytest.raises(DistributionWorkbookError, match="different source SPD"):
        load_distribution_targets(
            path2,
            rail_ids=("R1",),
            model_ids=("M1",),
            current_present={("R1", "M1"): 3},
            current_source_sha256="b" * 64,
        )


def test_format5_round_trips_signal_routing_policy_and_asset(tmp_path: Path) -> None:
    path = tmp_path / "protected.xlsx"
    asset_sha = "a" * 64
    content_sha = "b" * 64
    write_distribution_workbook(
        path,
        (),
        ("PWR NET", "M1\nPresent", "M1\nTarget", "M1\nTolerance (%)"),
        (("V1 (R1)", 1, 0, 0.0), ("V2 (R2)", 0, 1, 0.0)),
        metadata={
            **_current_metadata(),
            "Signal Routing Protection": "ON",
            "Routing Protection Scope": "SIGNAL_NET_ONLY",
            "Routing Policy Version": "SIGNAL_NET_ONLY_RESEARCH_V1",
            "Routing Clearance Mode": "FIXED_UM",
            "Trace-to-via Clearance (um)": 12.5,
            "Routing Asset SHA-256": asset_sha,
            "Routing Asset Content SHA-256": content_sha,
        },
    )

    imported = load_distribution_targets(
        path,
        rail_ids=("R1", "R2"),
        model_ids=("M1",),
        current_present={("R1", "M1"): 1, ("R2", "M1"): 0},
        current_routing_asset_sha256=asset_sha,
        current_routing_asset_content_sha256=content_sha,
    )

    assert imported.format_version == 5
    assert imported.routing_protection_enabled is True
    assert imported.routing_scope == "SIGNAL_NET_ONLY"
    assert imported.routing_clearance_um == 12.5
    assert imported.routing_asset_sha256 == asset_sha
    assert imported.routing_asset_content_sha256 == content_sha


def test_format5_protected_workbook_fails_closed_on_asset_mismatch(
    tmp_path: Path,
) -> None:
    path = tmp_path / "protected-mismatch.xlsx"
    write_distribution_workbook(
        path,
        (),
        ("PWR NET", "M1\nTarget"),
        (("V1 (R1)", 1),),
        metadata={
            **_current_metadata(),
            "Signal Routing Protection": "ON",
            "Routing Protection Scope": "SIGNAL_NET_ONLY",
            "Routing Policy Version": "SIGNAL_NET_ONLY_RESEARCH_V1",
            "Routing Clearance Mode": "FIXED_UM",
            "Trace-to-via Clearance (um)": 0.0,
            "Routing Asset SHA-256": "a" * 64,
            "Routing Asset Content SHA-256": "b" * 64,
        },
    )

    with pytest.raises(DistributionWorkbookError, match="routing asset differs"):
        load_distribution_targets(
            path,
            rail_ids=("R1",),
            model_ids=("M1",),
            current_present={("R1", "M1"): 1},
            current_routing_asset_sha256="c" * 64,
            current_routing_asset_content_sha256="b" * 64,
        )


def test_format5_protected_workbook_rejects_unknown_policy_version(
    tmp_path: Path,
) -> None:
    path = tmp_path / "protected-policy-mismatch.xlsx"
    write_distribution_workbook(
        path,
        (),
        ("PWR NET", "M1\nTarget"),
        (("V1 (R1)", 1),),
        metadata={
            **_current_metadata(),
            "Signal Routing Protection": "ON",
            "Routing Protection Scope": "SIGNAL_NET_ONLY",
            "Routing Policy Version": "SIGNAL_NET_ONLY_V999",
            "Routing Clearance Mode": "FIXED_UM",
            "Trace-to-via Clearance (um)": 0.0,
            "Routing Asset SHA-256": "a" * 64,
            "Routing Asset Content SHA-256": "b" * 64,
        },
    )

    with pytest.raises(DistributionWorkbookError, match="unsupported.*Policy Version"):
        load_distribution_targets(
            path,
            rail_ids=("R1",),
            model_ids=("M1",),
            current_present={("R1", "M1"): 1},
        )
